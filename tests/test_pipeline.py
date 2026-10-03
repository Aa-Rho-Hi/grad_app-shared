"""End-to-end + unit checks for the deterministic pipeline."""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import yaml
from src.normalize import build_record
from src.parse import read_rows, group_indexed
from src.validate import validate_record
from src.template import render
from src.run import run

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CFG = yaml.safe_load(open(os.path.join(ROOT, "config", "csv_field_map.yaml")))
SCHEMA = os.path.join(ROOT, "student.schema.json")
# Committed synthetic fixture (fake applicants) — data/ is gitignored, so tests
# must never depend on anything under it or they break on a clean CI checkout.
CSV = os.path.join(ROOT, "tests", "fixtures", "sample_applications.csv")


def _rec(cas_id):
    for r in read_rows(CSV):
        rec = build_record(r, CFG, "sample.csv", 0)
        if rec["cas_id"] == cas_id:
            return rec
    raise AssertionError(cas_id)


def test_indexed_grouping():
    _, groups = group_indexed({"ielts_overall_band_score_0": "7.0",
                               "first_name": "X",
                               "gpas_by_transcript_gpa_2": "3.5"})
    assert groups["ielts_overall_band_score"][0] == "7.0"
    assert groups["gpas_by_transcript_gpa"][2] == "3.5"


def test_ielts_attempts_and_best():
    rec = _rec("CAS1001")
    eng = rec["english_proficiency"]
    assert len(eng["ielts"]) == 2
    assert eng["best_ielts_overall"] == 7.5


def test_gpa_scale_detection():
    # 8.6 -> 10-point scale; 3.7 -> 4-point scale
    g1 = _rec("CAS1001")["education"][0]["gpa"]
    assert g1["scale"] == 10.0 and 3.4 <= g1["normalized_4"] <= 3.5
    g2 = _rec("CAS1002")["education"][0]["gpa"]
    assert g2["scale"] == 4.0 and g2["normalized_4"] == 3.7
    g3 = _rec("CAS1003")["education"][0]["gpa"]
    assert g3["scale"] == 100.0  # 72 -> percentage


def test_toefl_total_parsed():
    rec = _rec("CAS1002")
    assert rec["english_proficiency"]["best_toefl_total"] == 112.0


def test_summary_is_descriptive_not_scored():
    rec = _rec("CAS1001")
    text = render(rec)
    assert rec["personal"]["full_name"].split()[0] in text
    # neutral summary: no score/fit/recommendation language
    low = text.lower()
    assert "fit score" not in low and "/100" not in low
    assert "why consider" not in low and "strong candidate" not in low


def test_all_valid_against_schema():
    for r in read_rows(CSV):
        rec = build_record(r, CFG, "sample.csv", 0)
        ok, errors = validate_record(rec, SCHEMA)
        assert ok, errors


def test_run_builds_unified_student(tmp_path):
    from src.store import Store, db_path_for
    summary = run(CSV, str(tmp_path),
                  os.path.join(ROOT, "config", "csv_field_map.yaml"))
    assert summary["valid"] == 4 and summary["quarantined"] == 0
    store = Store(db_path_for(str(tmp_path)))
    u = store.get("CAS1001")
    assert u is not None
    # unified shape: identity + sectioned sources + provenance
    assert u["student_id"] == "CAS1001"
    assert u["sources"]["application"]["education"]
    assert u["sources"]["sop"] is None and u["sources"]["lors"] == []
    assert u["provenance"][0]["source"] == "application"
    assert u["summary"]
    assert store.count() == 4 and len(store.index()) == 4


def test_merge_accumulates_across_runs(tmp_path):
    from src.merge import upsert_source
    from src.store import Store, db_path_for
    # simulate: CSV run already wrote a student, then an SOP arrives
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    store = Store(db_path_for(str(tmp_path)))

    def _add_sop(u):
        upsert_source(u, "sop", {"text": "x", "word_count": 1}, file="s.txt")
        return u

    u = store.update("CAS1001", _add_sop)
    assert u["sources"]["application"] is not None     # CSV section preserved
    assert u["sources"]["sop"] == {"text": "x", "word_count": 1}  # new section added
    assert len(u["provenance"]) == 2
    # persisted, not just in memory
    assert store.get("CAS1001")["sources"]["sop"]["text"] == "x"


def test_sop_analysis_detects_areas_and_flags():
    from src.analyze import analyze_text
    a = analyze_text("My research is in VLSI and machine learning using Verilog and PyTorch. "
                     "I designed CMOS chips and neural networks.", "sop")
    assert "VLSI / Microelectronics" in a["detected_areas"]
    assert "Machine Learning / AI" in a["detected_areas"]
    assert "verilog" in a["mentioned_skills"] and "pytorch" in a["mentioned_skills"]
    # short SOP -> flagged
    assert any("very short" in f for f in a["flags"])


def test_document_analysis_extracts_review_profile():
    from src.analyze import analyze_text
    resume = (
        "Education\nMichigan State University\nBachelor of Science in Electrical Engineering\n"
        "Graduated with Honor GPA: 3.81\nGRE: 317 (V 148; Q 169)\n"
        "Skills: Python, Matlab, C++, SPICE, PCB design\n"
        "Research Experience\nOptimized a power electronics model and simulated circuits in Matlab.\n"
        "Course Project\nDeveloped a microcontroller project for signal acquisition.\n"
        "Won the department design award.\n"
    )
    rp = analyze_text(resume, "resume")["profile"]
    assert any("GPA: 3.81" in x for x in rp["education"])
    assert any("GRE: 317" in x for x in rp["test_scores"])
    assert rp["technical_preparation"]
    assert rp["experience"]
    assert rp["projects"]
    assert rp["distinctions"]

    sop = (
        "My goal is to research power electronics and improve grid reliability for the future. "
        "During my undergraduate laboratory research I simulated converter controls and learned "
        "measurement techniques that shaped my design intuition. "
        "The department faculty and research group at Texas A&M University align closely with my interests. "
        "Eventually I see myself contributing to the renewable energy industry and its growing workforce."
    )
    sp = analyze_text(sop, "sop")["profile"]
    assert sp["goals"] and sp["preparation"] and sp["program_fit"] and sp["career_direction"]

    lor = (
        "Dear Admissions Committee. I have known the applicant for two years in my research lab. "
        "He completed an independent project, showed strong analytical ability, and I strongly "
        "recommend him for graduate study."
    )
    lp = analyze_text(lor, "lor")["profile"]
    assert lp["evidence"]


def test_summary_surfaces_detailed_application_materials():
    """The scannable card still surfaces résumé, SOP and recommendation evidence
    for decision-making — clipped, under a heading, one item per line."""
    from src.documents import build_section
    resume = build_section("resume", (
        "Education\nMichigan State University\nBachelor of Science in Electrical Engineering\n"
        "Graduated with Honor GPA: 3.81\nGRE: 317 (V 148; Q 169)\n"
        "Skills: Python, Matlab, C++, SPICE, PCB design\n"
        "Research Experience\nOptimized a power electronics model and simulated circuits in Matlab.\n"
        "Course Project\nDeveloped a microcontroller project for signal acquisition.\n"
    ), "resume.txt")
    sop = build_section("sop", (
        "My goal is to study power systems at Texas A&M University. "
        "In my undergraduate research project, I simulated converter controls in Matlab and learned "
        "how laboratory measurements affect design decisions. "
        "I plan to use the graduate program to prepare for a career improving reliable energy systems."
    ), "sop.txt")
    lor = build_section("lor", (
        "Dear Admissions Committee. I have known the applicant for two years in my research lab. "
        "He completed an independent project, showed strong analytical ability, and I strongly "
        "recommend him for graduate study. Sincerely, Dr. Jane Smith."
    ), "lor.txt", recommender="Dr. Jane Smith")
    rec = {
        "cas_id": "CAS-RICH",
        "personal": {"full_name": "Detailed Applicant"},
        "_documents": {"resume": resume, "sop": sop, "scores": None, "lors": [lor]},
        "meta": {"warnings": []},
    }
    text = render(rec)
    assert "### Experience" in text and "power electronics model" in text
    assert "### In their words (SOP)" in text and "Texas A&M University" in text
    assert "### Letters" in text and "Dr. Jane Smith" in text
    assert "### On file" in text
    low = text.lower()
    assert "fit score" not in low and "/100" not in low and "strong candidate" not in low
    # Bounded, not terse: the card carries whole claims now (the clip budgets
    # were raised once bullets stopped being cut mid-sentence), but it is still
    # a scan aid rather than a transcript.
    assert len(text.split()) < 400, "summary card drifted back into essay length"
    # Readability contract: every section is a heading plus one item per line,
    # so nothing is a semicolon-joined paragraph a reviewer has to unpack.
    for line in text.splitlines():
        assert len(line) < 420, f"line too long to scan: {line[:80]}…"
    for block in text.split("\n\n")[1:]:
        head, *items = block.splitlines()
        assert head.startswith("### "), f"section without a heading: {head[:60]}"
        assert items and all(i.startswith("- ") for i in items), f"unbulleted: {head}"


def test_reindex_refreshes_existing_summaries(tmp_path):
    from src.documents import build_section
    from src.merge import new_student, summary_view, upsert_source
    from src.run import reindex
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    u = new_student("CAS-REFRESH")
    section = build_section("resume", (
        "Education\nState University\nBachelor of Science in Electrical Engineering\n"
        "GPA: 3.7\nSkills: Python and Matlab\nResearch Experience\n"
        "Developed a power systems simulation project."
    ), "resume.txt")
    section["analysis"].pop("profile")                 # simulate a pre-upgrade stored record
    upsert_source(u, "resume", section, file="resume.txt")
    u["summary"] = "OLD SUMMARY"
    store.put(u)
    reindex(str(tmp_path))
    refreshed = Store(db_path_for(str(tmp_path))).get("CAS-REFRESH")
    assert "OLD SUMMARY" not in refreshed["summary"]
    assert "### Experience" in refreshed["summary"]
    assert render(summary_view(refreshed)) == refreshed["summary"]


def test_duplicate_sop_detection(tmp_path):
    from src.store import Store, db_path_for
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    essay = ("I am deeply passionate about power systems and renewable energy. "
             "My goal is to research smart grids and advance the field. ") * 8
    ingest_document_txt(tmp_path, "CAS1001", essay)
    ingest_document_txt(tmp_path, "CAS1002", essay)   # same essay -> should flag CAS1002
    u = Store(db_path_for(str(tmp_path))).get("CAS1002")
    flags = u["sources"]["sop"]["analysis"]["flags"]
    assert any("similar" in f.lower() for f in flags)


def ingest_document_txt(tmp_path, cas_id, text):
    from src.documents import ingest_document
    p = tmp_path / f"{cas_id}_sop.txt"
    p.write_text(text)
    ingest_document(str(p), cas_id, "sop", str(tmp_path))


def test_document_ingest_merges_sop(tmp_path):
    from src.documents import ingest_document
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    sop = tmp_path / "CAS1001_sop.txt"
    sop.write_text("I am passionate about power systems and renewable energy research. " * 20)
    unified, words = ingest_document(str(sop), "CAS1001", "sop", str(tmp_path))
    assert words > 50
    assert unified["sources"]["application"] is not None      # CSV section kept
    assert unified["sources"]["sop"]["word_count"] == words    # SOP section added
    low = unified["summary"].lower()
    assert "in their words (sop)" in low and "- goal:" in low  # SOP surfaced in the card
    # provenance records both sources
    assert {p["source"] for p in unified["provenance"]} == {"application", "sop"}


def test_reconcile_merges_email_and_cas_records(tmp_path):
    from src.merge import new_student, upsert_source
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    # same person under two keys: email (CSV) and numeric cas_id (PDF)
    a = new_student("jane@x.com"); a["identity"]["full_name"] = "Jane Doe"
    upsert_source(a, "application", {"programs": [{"name": "MS EE"}]}, file="app.csv")
    b = new_student("1000123456"); b["identity"]["full_name"] = "Jane Doe"
    upsert_source(b, "sop", {"text": "hi", "word_count": 1}, file="x.pdf")
    store.put(a); store.put(b)
    merged = store.reconcile()
    assert merged == 1
    ids = [u["student_id"] for u in store.all()]
    assert ids == ["1000123456"]                             # kept the cas_id key
    u = store.get("1000123456")
    assert u["sources"]["application"] and u["sources"]["sop"]  # both sources merged


def test_concurrent_updates_do_not_lose_data(tmp_path):
    """Two writers hammering the same cas_id must not clobber each other —
    every append must survive (no lost update)."""
    import threading
    from src.merge import new_student
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    store.put(new_student("CAS9999"))
    N = 40

    def append(i):
        store.update("CAS9999",
                     lambda u: u["warnings"].append(f"note-{i}") or u)

    threads = [threading.Thread(target=append, args=(i,)) for i in range(N)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    u = store.get("CAS9999")
    assert len(u["warnings"]) == N                            # all writes survived
    assert len(set(u["warnings"])) == N


def test_ocr_reads_scanned_pdf():
    import pytest, io
    pytest.importorskip("PIL")
    pytest.importorskip("pypdfium2")
    from src import ocr
    if not ocr.available():
        pytest.skip("tesseract not installed")
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1200, 400), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf", 44)
    except Exception:
        font = ImageFont.load_default()
    d.text((40, 120), "power systems and machine learning", fill="black", font=font)
    buf = io.BytesIO(); img.save(buf, format="PDF")
    pages, info = ocr.pdf_page_texts(buf.getvalue())
    assert info["ocr_used"] == 1
    assert "power systems" in pages[0].lower()


def test_cas_id_and_name_from_filename():
    from src.packet import cas_id_from_filename, name_from_filename
    fn = "1000419138_Joonha_Jun_Full Application_200948_Fall 2020 M.S. EE.pdf"
    assert cas_id_from_filename(fn) == "1000419138"
    nm = name_from_filename(fn)
    assert nm["first_name"] == "Joonha" and nm["last_name"] == "Jun"


def _make_packet(path, name, sop_text, resume_text):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import LETTER
    c = canvas.Canvas(path, pagesize=LETTER)
    for title, body in [("Application for Admission", f"Applicant: {name}"),
                        ("Statement of Purpose", sop_text),
                        ("Curriculum Vitae", resume_text),
                        ("IELTS Test Report Form", "Overall Band Score 7.5")]:
        c.setFont("Helvetica-Bold", 16); c.drawString(72, 720, title)
        c.setFont("Helvetica", 11); c.drawString(72, 690, body)
        c.showPage()
    c.save()


def test_zip_ingestion_segments_and_merges(tmp_path):
    import pytest, zipfile
    pytest.importorskip("reportlab")
    from src.ingest_zip import ingest_zip
    from src.store import Store, db_path_for
    pdf = tmp_path / "1000419138_Joonha_Jun_Full Application_1_Fall 2020 EE.pdf"
    _make_packet(str(pdf), "Joonha Jun",
                 "My research is in power systems and smart grid renewable energy. " * 10,
                 "Skills: python, matlab, verilog. Experience: research intern. " * 5)
    zpath = tmp_path / "packets.zip"
    with zipfile.ZipFile(zpath, "w") as z:
        z.write(pdf, pdf.name)
    report = ingest_zip(str(zpath), str(tmp_path))
    assert report["pdfs_found"] == 1 and report["students_updated"] == 1
    u = Store(db_path_for(str(tmp_path))).get("1000419138")
    assert u["sources"]["sop"] and u["sources"]["resume"]          # both kept
    assert "Power Systems" in u["sources"]["sop"]["analysis"]["detected_areas"]
    assert "verilog" in u["sources"]["resume"]["analysis"]["mentioned_skills"]
    assert u["identity"]["full_name"] == "Joonha Jun"              # name from filename
    assert u["sources"]["scores"]                                  # IELTS page kept as scores


def test_extract_recommender_and_scores():
    from src.analyze import extract_recommender, extract_test_scores
    lor = ("REFERENCES\nSarang Dhongdi\nProfessional Title: Assistant Professor\n"
           "Organization: BITS Pilani K K Birla Goa Campus\nEmail: x@y.edu\n")
    rec = extract_recommender(lor)
    assert rec["name"] == "Sarang Dhongdi"
    assert rec["title"] == "Assistant Professor"
    assert "BITS Pilani" in rec["organization"]
    gre = ("verbal reasoning with your scaled score of 155 out of 170. "
           "quantitative reasoning with your scaled score of 168 out of 170. "
           "analytical writing with your score of 4.5 out of 6.")
    ts = extract_test_scores(gre)
    assert ts["gre_verbal"] == 155 and ts["gre_quant"] == 168
    assert ts["gre_awa"] == 4.5 and ts["gre_total"] == 323
    assert extract_test_scores("Overall Band Score 7.5")["ielts_overall"] == 7.5
    assert extract_recommender("no references here at all") == {}


# --------------------------------------------------------------------------
# Edge cases
# --------------------------------------------------------------------------

def _write_csv(path, header, *rows):
    with open(path, "w", encoding="utf-8") as f:
        f.write(header + "\n")
        for r in rows:
            f.write(r + "\n")


def test_invalid_row_is_quarantined(tmp_path):
    """A value outside the allowed range (IELTS > 9) must be quarantined with a
    readable reason — not silently coerced or dropped — while valid rows survive."""
    from src.store import Store, db_path_for
    csv = tmp_path / "mixed.csv"
    _write_csv(str(csv),
               "cas_id,first_name,last_name,ielts_overall_band_score_0",
               "CAS5001,Bad,Row,11.0",       # 11.0 > 9 -> invalid
               "CAS5002,Good,Row,7.0")        # fine
    field_map = os.path.join(ROOT, "config", "csv_field_map.yaml")
    summary = run(str(csv), str(tmp_path), field_map)
    assert summary["valid"] == 1 and summary["quarantined"] == 1
    store = Store(db_path_for(str(tmp_path)))
    assert store.get("CAS5002") is not None and store.get("CAS5001") is None
    q = store.quarantine()
    assert len(q) == 1 and q[0]["cas_id"] == "CAS5001"
    assert any("less than or equal to 9" in e for e in q[0]["errors"])


def test_corrupt_pdf_in_zip_is_reported_not_fatal(tmp_path):
    """A non-PDF / corrupt file in the ZIP is reported as unmatched and does not
    crash the batch or create a record."""
    import zipfile
    from src.ingest_zip import ingest_zip
    from src.store import Store, db_path_for
    zpath = tmp_path / "packets.zip"
    with zipfile.ZipFile(zpath, "w") as z:
        z.writestr("1000222222_Bad_File_Full Application_x.pdf", b"this is not a pdf")
    report = ingest_zip(str(zpath), str(tmp_path))
    assert report["pdfs_found"] == 1
    assert report["students_updated"] == 0
    assert report["unmatched_files"]                       # reported, not raised
    assert Store(db_path_for(str(tmp_path))).count() == 0


def test_duplicate_cas_id_csv_then_zip_merges_one_record(tmp_path):
    """The same numeric cas_id arriving from the application path and from a PDF
    packet must end up as ONE record holding both sources — no duplicate."""
    import pytest, zipfile
    pytest.importorskip("reportlab")
    from src.merge import new_student, upsert_source
    from src.ingest_zip import ingest_zip
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    # application section already on disk for this numeric id
    a = new_student("1000419138")
    a["identity"]["full_name"] = "Joonha Jun"
    upsert_source(a, "application", {"programs": [{"name": "MS EE"}]}, file="app.csv")
    store.put(a)
    # now a packet for the SAME id arrives
    pdf = tmp_path / "1000419138_Joonha_Jun_Full Application_1_Fall 2020 EE.pdf"
    _make_packet(str(pdf), "Joonha Jun",
                 "My research is in power systems and smart grid renewable energy. " * 10,
                 "Skills: python, matlab, verilog. Experience: research intern. " * 5)
    zpath = tmp_path / "p.zip"
    with zipfile.ZipFile(zpath, "w") as z:
        z.write(pdf, pdf.name)
    ingest_zip(str(zpath), str(tmp_path))
    assert store.count() == 1                               # not duplicated
    u = store.get("1000419138")
    assert u["sources"]["application"] and u["sources"]["sop"] and u["sources"]["resume"]


def _make_headingless_packet(path):
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import LETTER
    c = canvas.Canvas(path, pagesize=LETTER)
    # page 1 — resume with NO heading, just a contact block
    lines1 = ["Jane Doe", "jane.doe@example.com  +1 415 555 1234",
              "San Francisco, CA", "Skills: python, matlab, verilog",
              "Experience: research intern at a power lab"]
    y = 720
    for ln in lines1:
        c.setFont("Helvetica", 12); c.drawString(72, y, ln); y -= 20
    c.showPage()
    # page 2 — statement with NO heading, just prose full of SOP cues
    sop = ("My research interest is in power systems and renewable energy. "
           "I am applying to the masters program because my goal is to study smart grids. "
           "I plan to pursue a master degree and my motivation is to advance the field. "
           "In my undergraduate I built models and my interest in control grew. "
           "I would love to research grid stability and I intend to contribute to the area. ") * 2
    y = 720
    for chunk in [sop[i:i+90] for i in range(0, len(sop), 90)]:
        c.setFont("Helvetica", 11); c.drawString(60, y, chunk); y -= 16
    c.showPage()
    c.save()


def test_headingless_packet_segments_by_heuristics(tmp_path):
    """A packet with no section headings — resume is just a contact block, the
    statement is just prose — is still recovered via the content heuristics."""
    import pytest
    pytest.importorskip("reportlab")
    from src.packet import segment
    pdf = tmp_path / "headingless.pdf"
    _make_headingless_packet(str(pdf))
    seg = segment(str(pdf))
    assert "resume" in seg["sections"]                     # recovered by contact block
    assert "sop" in seg["sections"]                        # recovered by prose + cues


def test_mixed_gpa_scales_normalize_to_4(tmp_path):
    """GPAs on different university scales each auto-detect and normalize to /4.0."""
    from src.normalize import _normalize_gpa
    cases = {4.0: 3.9, 5.0: 4.6, 10.0: 8.6, 20.0: 18.0, 100.0: 72.0}
    for expected_scale, raw in cases.items():
        g = _normalize_gpa(raw, CFG)
        assert g["scale"] == expected_scale, (raw, g)
        assert 0.0 <= g["normalized_4"] <= 4.0
    # a 10-scale 8.6 lands around 3.44/4.0; a 4-scale 3.9 stays ~3.9
    assert abs(_normalize_gpa(8.6, CFG)["normalized_4"] - 3.44) < 0.05
    assert abs(_normalize_gpa(3.9, CFG)["normalized_4"] - 3.9) < 0.05


def test_store_works_without_os_locks(tmp_path, monkeypatch):
    """On a filesystem with no POSIX locks (some network/FUSE mounts), the store
    falls back to nolock + in-process serialization and still reads/writes/persists."""
    from src.store import Store, db_path_for
    from src.merge import new_student
    # force the 'filesystem can't lock' branch
    monkeypatch.setattr(Store, "_fs_supports_locks", lambda self: False)
    s = Store(db_path_for(str(tmp_path)))
    assert s._cfg["nolock"] is True and s._cfg["journal"] == "off"
    s.put(new_student("X1"))
    s.update("X1", lambda u: u["warnings"].append("w") or u)
    rec = s.get("X1")
    assert rec is not None and rec["warnings"] == ["w"]
    assert s.count() == 1
    # a second Store on the same path sees the persisted record
    assert Store(db_path_for(str(tmp_path))).get("X1")["warnings"] == ["w"]


def test_ocr_reads_skewed_scanned_page(tmp_path):
    """A skewed image-only page is deskewed + preprocessed and still read."""
    import pytest, io
    pytest.importorskip("PIL")
    pytest.importorskip("pypdfium2")
    pytest.importorskip("numpy")
    from src import ocr
    if not ocr.available():
        pytest.skip("tesseract not installed")
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (1600, 500), "white")
    d = ImageDraw.Draw(img)
    try:
        font = ImageFont.truetype("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", 38)
    except Exception:
        font = ImageFont.load_default()
    d.text((40, 160), "power systems and machine learning", fill="black", font=font)
    img = img.rotate(3, resample=Image.BILINEAR, fillcolor="white", expand=True)
    buf = io.BytesIO(); img.save(buf, format="PDF")
    pages, info = ocr.pdf_page_texts(buf.getvalue())
    assert info["ocr_used"] == 1
    low = pages[0].lower()
    assert "power systems" in low and "machine learning" in low


def test_clean_ocr_text_fixes_common_artifacts():
    from src.ocr import clean_ocr_text
    assert clean_ocr_text("| have known Lekhaj") == "I have known Lekhaj"
    assert clean_ocr_text("ideas, | was able to observe") == "ideas, I was able to observe"
    assert clean_ocr_text("|n my view |t is fine") == "In my view It is fine"
    assert clean_ocr_text("the efﬁcient ﬂow is ﬁne") == "the efficient flow is fine"
    assert clean_ocr_text("collapse    spaces") == "collapse spaces"
    assert clean_ocr_text("code a||b stays") == "code a||b stays"   # not standalone


def test_otsu_binarize_is_deterministic():
    import pytest
    pytest.importorskip("numpy")
    from PIL import Image
    import numpy as np
    from src.ocr import _otsu_binarize
    arr = np.zeros((20, 20), dtype=np.uint8)
    arr[5:15, 5:15] = 200                      # a bright square on dark
    img = Image.fromarray(arr, mode="L")
    a = np.asarray(_otsu_binarize(img))
    b = np.asarray(_otsu_binarize(img))
    assert (a == b).all()                      # deterministic
    assert set(np.unique(a)).issubset({0, 255})


def test_extract_standardized_tests_cas_form():
    """The CAS academic-history GRE/IELTS/TOEFL tables (date row with interleaved
    percentiles) parse to the right figures."""
    from src.analyze import extract_standardized_tests
    text = (
        "STANDARDIZED TESTS\n"
        "OFFICIAL GRE\n"
        "Date ETS Registration Code Verbal Quantitative Analytical Writing\n"
        "09-20-2019 148 37% 164 84% 3.5 39%\n"
        "UNOFFICIAL IELTS\n"
        "Date Candidate Number TRF Number Listening Reading Speaking Writing Overall Band\n"
        "02-02-2019 012440 8.5 7.5 6.0 6.0 7.0\n"
        "OFFICIAL TOEFL\n"
        "Date Type Listening Reading Speaking Writing Essay Total\n"
        "09-28-2019 Internet-based 23 21 23 27 94\n"
    )
    ts = extract_standardized_tests(text)
    assert ts["gre_verbal"] == 148 and ts["gre_quant"] == 164 and ts["gre_awa"] == 3.5
    assert ts["gre_total"] == 312
    assert ts["ielts_overall"] == 7.0
    assert ts["toefl_total"] == 94


def test_likert_column_anchors_and_nearest():
    """The rating columns are found from evenly-spaced (1)-(5) markers, and a
    phone number / date that merely contains those digits is rejected."""
    from src.ratings import _column_anchors, _nearest_rating
    header = [(309.0, "1"), (358.0, "2"), (430.0, "3"), (509.0, "4"), (585.0, "5"),
              (120.0, "A"), (250.0, "x")]
    rows = {512: header}
    anchors, y = _column_anchors(rows)
    assert y == 512 and anchors == {1: 309.0, 2: 358.0, 3: 430.0, 4: 509.0, 5: 585.0}
    # checkmark x -> rating column
    assert _nearest_rating(561, anchors) == 5     # the real packet's marks
    assert _nearest_rating(483, anchors) == 4
    assert _nearest_rating(309, anchors) == 1
    # a phone number row "821030756824" -> digits not strictly increasing/even
    phone = [(float(100 + i * 7), ch) for i, ch in enumerate("821030756824")]
    assert _column_anchors({400: phone}) == (None, None)


def test_likert_name_and_recommendation_parsing():
    from src.ratings import _evaluator_name, _recommendation
    txt = ("EVALUATIONS\nEVALUATOR INFORMATION\nWongyu Bae\nTitle: Assistant Professor\n"
           "REFERENCE RATINGS\nI highly recommend this applicant\n")
    assert _evaluator_name(txt) == "Wongyu Bae"
    assert _recommendation(txt) == "i highly recommend"
    assert _recommendation("no rec phrase here") is None


def test_segmenter_labels_pages_by_content_not_neighbour():
    """The per-page scorer + Viterbi must label each page by its own content, so
    a thesis page or a duplicate resume can't bleed into an adjacent section."""
    from src.packet import _page_scores, _viterbi
    headers = CFG and __import__("yaml").safe_load(
        open(os.path.join(ROOT, "config", "keywords.yaml")))["section_headers"]
    pages = [
        "Application for Admission\nApplicant: X\nProgram Level: Masters\nStart Term: Fall",
        "John Doe\njohn@example.com  +1 415 555 1234\nEDUCATION\nEXPERIENCE\nSKILLS: python, matlab",
        ("Statement of Purpose\nMy research interest is in power systems. I am applying "
         "because my goal is to study smart grids. I plan to pursue research and my "
         "motivation is to advance the field. " * 3),
        ("Type: Letter of Reference and Likert Scale\nREFERENCES\nDr. Smith\n"
         "Professional Title: Professor\nWaiver of Evaluation: Yes\nI strongly recommend the applicant."),
        "Test Taker Score Report\nTOEFL\nVerbal Reasoning scaled score 160, 88th percentile\nTest Date: 2019",
        "Chapter 2\nMathematical Background\nBibliography\nList of Figures\n" + "∑ α β ∫ ≥ θ λ " * 5,
        "Jane Roe\njane@example.com  +1 212 555 9000\nWORK EXPERIENCE\nPROJECTS\nSKILLS: c++, verilog",
    ]
    labels = _viterbi([_page_scores(p, headers) for p in pages])
    assert labels[0] == "ignore"    # application form
    assert labels[1] == "resume"
    assert labels[2] == "sop"
    assert labels[3] == "lor"
    assert labels[4] == "scores"
    assert labels[5] == "ignore"    # thesis page dropped, NOT merged into scores/sop
    assert labels[6] == "resume"    # a resume-like page after others stays resume


def test_segmenter_distinguishes_recommendation_letter_from_sop():
    """A narrative recommendation letter is first-person too — it must still be
    labelled lor (salutation + third-person about the applicant + recommend),
    not sop."""
    from src.packet import _page_scores, _viterbi
    headers = __import__("yaml").safe_load(
        open(os.path.join(ROOT, "config", "keywords.yaml")))["section_headers"]
    rec_letter = ("Dear Members of the Admission Committee, My name is Professor Bae. "
                  "I have known Mr. Jun for three years as his research advisor. He is an "
                  "outstanding student and his work in his lab was excellent. He led his "
                  "team and his projects with great skill. I have seen him grow. I highly "
                  "recommend him for your program. Yours sincerely, Professor Bae. ") * 2
    sop = ("Personal Statement. My research interest is in power systems and control. "
           "I decided to pursue a master's degree because my goal is to study smart grids. "
           "I plan to research grid stability and my motivation is to advance the field. "
           "I came to this decision after my undergraduate projects. ") * 2
    labels = _viterbi([_page_scores(rec_letter, headers), _page_scores(sop, headers)])
    assert labels[0] == "lor"
    assert labels[1] == "sop"


def test_segmenter_detects_essay_style_sop():
    """An SOP that opens with a third-person hook (no 'I am/my goal' up front) is
    still detected via first-person density + SOP vocabulary, not mislabelled
    résumé."""
    from src.packet import _page_scores, _viterbi
    headers = __import__("yaml").safe_load(
        open(os.path.join(ROOT, "config", "keywords.yaml")))["section_headers"]
    essay_sop = ("With great power comes great electricity bill. More than a billion "
                 "people still lack access to reliable power across the world. ") + (
                 "I worked on smart-grid projects and I learned a great deal about control. "
                 "I want to pursue graduate study and my research interest is power systems. "
                 "My motivation drives me toward a doctoral career and admission to this program. ") * 4
    resume_like = ("Jane Roe  jane@example.com  +1 415 555 1234  EDUCATION  WORK EXPERIENCE  "
                   "PROJECTS  SKILLS: python, matlab, verilog ") * 3
    labels = _viterbi([_page_scores(resume_like, headers), _page_scores(essay_sop, headers)])
    assert labels[0] == "resume"
    assert labels[1] == "sop"


def test_empty_csv_rows_are_quarantined(tmp_path):
    """Blank / unmapped CSV rows (no cas_id, email, name, or data) are quarantined
    rather than surfaced as empty ROW-n applicants."""
    from src.store import Store, db_path_for
    csv = tmp_path / "rows.csv"
    csv.write_text("cas_id,first_name,last_name,ielts_overall_band_score_0\n"
                   "CAS900,Real,Person,7.0\n,,,\n,,,\n")
    field_map = os.path.join(ROOT, "config", "csv_field_map.yaml")
    summary = run(str(csv), str(tmp_path), field_map)
    assert summary["valid"] == 1 and summary["quarantined"] == 2
    store = Store(db_path_for(str(tmp_path)))
    assert [e["cas_id"] for e in store.index()] == ["CAS900"]
    assert all(c["cas_id"].startswith("ROW-") for c in store.quarantine())


def test_needs_ocr_catches_watermarked_scans():
    """Regression: a scanned page whose only embedded text is a watermark
    (e.g. 'Scanned by CamScanner', 21 chars) used to defeat the 20-char OCR
    threshold and be silently dropped. It must now be flagged for OCR."""
    from src.ocr import _needs_ocr
    assert _needs_ocr("Scanned by CamScanner", has_image=True) is True   # the bug case
    assert _needs_ocr("page 3", has_image=False) is True                 # near-empty
    dense = "word " * 200                                                # real text page
    assert _needs_ocr(dense, has_image=True) is False                    # don't OCR good text
    assert _needs_ocr(dense, has_image=False) is False


def test_page_has_image_detection():
    import io
    import pytest
    pytest.importorskip("reportlab")
    from PIL import Image
    from pypdf import PdfReader
    from src.ocr import _page_has_image
    # image-only PDF page -> detected as having an image
    im = Image.new("RGB", (200, 80), "white")
    buf = io.BytesIO(); im.save(buf, format="PDF")
    assert _page_has_image(PdfReader(io.BytesIO(buf.getvalue())).pages[0]) is True
    # a pure-text page -> no image
    from reportlab.pdfgen import canvas
    cb = io.BytesIO(); c = canvas.Canvas(cb)
    c.drawString(72, 720, "hello text only"); c.showPage(); c.save()
    assert _page_has_image(PdfReader(io.BytesIO(cb.getvalue())).pages[0]) is False


# --------------------------------------------------------------------------
# Web layer (Flask) — smoke tests through the test client
# --------------------------------------------------------------------------

def _make_group(name="Systems", leader="gratz-paul", members=("fan-zhiwen",)):
    """One review group in the per-test groups file (conftest isolates it)."""
    from src import groups as groupsmod
    from src.triage import load_professors
    groups = groupsmod.add_group(groupsmod.load_groups(), name, leader=leader,
                                 members=list(members), professors=load_professors())
    groupsmod.save_groups(groups)
    return groups[-1]["id"]


def _client(tmp_path, monkeypatch):
    from src import app as appmod
    monkeypatch.setattr(appmod, "DATA", str(tmp_path))
    os.makedirs(os.path.join(str(tmp_path), "raw"), exist_ok=True)
    appmod.app.config.update(TESTING=True)
    return appmod, appmod.app.test_client()


def test_web_upload_browse_download_reset(tmp_path, monkeypatch):
    import io
    appmod, client = _client(tmp_path, monkeypatch)

    # empty state renders
    assert client.get("/").status_code == 200

    # upload the sample CSV through the form
    with open(CSV, "rb") as f:
        data = {"csv": (io.BytesIO(f.read()), "sample_applications.csv")}
    r = client.post("/upload", data=data, content_type="multipart/form-data",
                    follow_redirects=True)
    assert r.status_code == 200

    # the store now has students; the list + a detail page render
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    assert store.count() == 4
    cas = store.index()[0]["cas_id"]
    assert client.get(f"/student/{cas}").status_code == 200
    assert client.get("/student/NOPE").status_code == 404

    # summaries page + report download
    assert client.get("/summaries").status_code == 200
    dl = client.get("/download/summaries")
    assert dl.status_code == 200 and b"summaries" in dl.data.lower()

    # reset clears the store
    client.post("/reset", follow_redirects=True)
    assert Store(db_path_for(str(tmp_path))).count() == 0


def test_the_four_steps_walk_in_order(tmp_path, monkeypatch):
    """Step 1 upload -> step 2 sheet -> step 3 marked-up sheet -> step 4 deck.

    The thing this guards: no summary card is shown before the status pass.
    """
    import io
    import openpyxl
    from src.review import STATUS_COLUMN
    from src.store import Store, db_path_for
    appmod, client = _client(tmp_path, monkeypatch)

    # --- step 1: upload only. No cards. ('/' is the dashboard now.) ---
    first = client.get("/step/1").get_data(as_text=True)
    assert "Step 1" in first and 'class="summary"' not in first
    with open(CSV, "rb") as f:
        r = client.post("/process", data={"csv": (io.BytesIO(f.read()), "roster.csv")},
                        content_type="multipart/form-data")
    assert r.headers["Location"].endswith("/step/2")      # straight on to the sheet
    store = Store(db_path_for(str(tmp_path)))
    assert store.count() == 4
    assert 'class="summary"' not in client.get("/step/1").get_data(as_text=True)

    # --- step 2: the sheet, still no cards ---
    two = client.get("/step/2")
    assert two.status_code == 200 and 'class="summary"' not in two.get_data(as_text=True)
    sheet = client.get("/review-sheet")
    assert sheet.status_code == 200

    # --- step 3: mark it up and send it back ---
    assert client.get("/step/3").status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(sheet.data))
    ws = wb["Applicants"]
    col = [c.value for c in ws[1]].index(STATUS_COLUMN) + 1
    for row in range(2, ws.max_row + 1):
        ws.cell(row=row, column=col).value = "tbd" if row < 4 else "accept"
    buf = io.BytesIO()
    wb.save(buf)
    r = client.post("/step/3", data={"csv": (io.BytesIO(buf.getvalue()), "roster_review.xlsx")},
                    content_type="multipart/form-data")
    assert r.headers["Location"].endswith("/triage")      # step 4, automatically

    # --- step 4: only the tbd ones are in the deck ---
    deck_page = client.get("/triage").get_data(as_text=True)
    assert deck_page.count('class="tcard"') == 2
    index = Store(db_path_for(str(tmp_path))).index()
    assert sum(1 for e in index if e["review_status"] == "tbd") == 2
    assert sum(1 for e in index if e["review_status"] == "accept") == 2

    # the browse view still exists, off the flow
    assert client.get("/applicants").status_code == 200


def test_upload_by_fetch_reports_progress_then_finishes(tmp_path, monkeypatch):
    """The browser posts by fetch and polls; a plain form post still works."""
    import io
    import time
    appmod, client = _client(tmp_path, monkeypatch)
    with open(CSV, "rb") as f:
        r = client.post("/process", data={"csv": (io.BytesIO(f.read()), "roster.csv")},
                        content_type="multipart/form-data",
                        headers={"X-Requested-With": "fetch"})
    assert r.status_code == 200
    job_id = r.get_json()["job"]

    for _ in range(200):                      # the fixture is small; ~instant
        status = client.get(f"/jobs/{job_id}").get_json()
        if status["done"]:
            break
        time.sleep(0.05)
    assert status["done"] and not status["error"]
    assert status["percent"] == 100
    assert "summarized" in (status["message"] or "")
    assert status["next"].endswith("/step/2")

    from src.store import Store, db_path_for
    assert Store(db_path_for(str(tmp_path))).count() == 4
    assert client.get("/jobs/nope").status_code == 404


def test_actions_answer_fetch_without_a_reload(tmp_path, monkeypatch):
    """Send-to-group, professor edits and clear-all return JSON to a fetch
    caller and keep redirecting for a plain form post."""
    appmod, client = _client(tmp_path, monkeypatch)
    gid = _make_group()
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    cas = store.index()[0]["cas_id"]
    store.update(cas, lambda u: {**u, "review_status": "tbd"})
    client.post("/triage/swipe", data={"cas_id": cas, "direction": "shortlist"})

    fetch = {"X-Requested-With": "fetch"}
    r = client.post("/triage/assign", data={"cas_id": cas, "group": gid}, headers=fetch)
    assert r.status_code == 200 and r.get_json()["name"] == "Systems"

    # a plain post still redirects, so the page works with JS off
    assert client.post("/triage/assign",
                       data={"cas_id": cas, "group": ""}).status_code == 302

    r = client.post("/reset", headers=fetch)
    assert r.status_code == 200 and r.get_json()["cleared"] == 4


def test_step_three_needs_a_spreadsheet(tmp_path, monkeypatch):
    appmod, client = _client(tmp_path, monkeypatch)
    r = client.post("/step/3", data={}, content_type="multipart/form-data",
                    follow_redirects=True)
    assert r.status_code == 200 and b"status column filled in" in r.data


def test_web_process_accepts_loose_packet_pdfs(tmp_path, monkeypatch):
    """The packets drop zone takes the PDFs themselves, not only a ZIP."""
    import io
    import pytest
    pytest.importorskip("reportlab")
    appmod, client = _client(tmp_path, monkeypatch)
    pdf = tmp_path / "1000419138_Joonha_Jun_Full Application_1_Fall 2020 EE.pdf"
    _make_packet(str(pdf), "Joonha Jun",
                 "My goal is to research power systems at Texas A&M University. " * 6,
                 "Education: State University. Bachelor of Science in Electrical Engineering.")
    r = client.post("/process", data={
        "zip": [(io.BytesIO(pdf.read_bytes()), pdf.name)],
    }, content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200
    from src.store import Store, db_path_for
    assert Store(db_path_for(str(tmp_path))).get("1000419138")["sources"]["sop"]


def test_web_process_rejects_other_packet_types(tmp_path, monkeypatch):
    import io
    appmod, client = _client(tmp_path, monkeypatch)
    r = client.post("/process", data={
        "zip": [(io.BytesIO(b"not a packet"), "notes.txt")],
    }, content_type="multipart/form-data", follow_redirects=True)
    assert r.status_code == 200
    assert b".zip or .pdf" in r.data


def test_web_upload_rejects_empty(tmp_path, monkeypatch):
    appmod, client = _client(tmp_path, monkeypatch)
    r = client.post("/upload", data={}, content_type="multipart/form-data",
                    follow_redirects=True)
    assert r.status_code == 200          # flashes an error, doesn't crash


# ---------------------------------------------------------- the review pass --

def _packet_record(resume_text="", scores_text="", status=None):
    from src.documents import build_section
    from src.merge import new_student, upsert_source
    u = new_student("CAS-REVIEW")
    if resume_text:
        upsert_source(u, "resume", build_section("resume", resume_text, "r.txt"), file="r.txt")
    if scores_text:
        upsert_source(u, "scores", build_section("scores", scores_text, "s.txt"), file="s.txt")
    if status:
        u["review_status"] = status
    return u


def test_same_name_under_two_ids_is_flagged_not_merged():
    """The roster and a packet filename can spell a name differently and carry
    different ids; the store then holds the same person twice."""
    from src.merge import _norm_name, likely_duplicate_groups
    assert _norm_name("CHEN-NI LIU") == _norm_name("Chen Ni Liu")   # was not equal
    groups = likely_duplicate_groups([
        {"cas_id": "1000419138", "name": "CHEN-NI LIU"},
        {"cas_id": "8929276570", "name": "Chen Ni Liu"},
        {"cas_id": "3", "name": "Someone Else"},
    ])
    assert [[m["cas_id"] for m in g] for g in groups] == [["1000419138", "8929276570"]]
    # One record per id is not a duplicate, however many sources it has.
    assert likely_duplicate_groups([{"cas_id": "1", "name": "A B"},
                                    {"cas_id": "1", "name": "A B"}]) == []


def test_gpa_shows_the_number_the_sheet_holds():
    """The export's own figure, not a conversion. A /4.0 denominator only
    appears when the sheet stated the scale — printing "3.8/4.0" for a value
    that might be out of 5 asserts something the data doesn't support."""
    warn = ("GPA scale was auto-detected from the value (e.g. 3.8 is read as "
            "/4.0 even if the transcript uses /5.0) — verify.")
    rec = {"cas_id": "G1", "personal": {"full_name": "Gpa Person"},
           "education": [{"college_name": "X",
                          "gpa": {"raw": 3.8, "scale": 4.0, "normalized_4": 3.8}}],
           "_documents": {"resume": None, "sop": None, "scores": None, "lors": []},
           "meta": {"warnings": [warn]}}
    text = render(rec)
    assert "GPA 3.8 (scale not stated in the sheet)" in text
    assert "3.8/4.0" not in text

    rec["meta"]["warnings"] = []                 # scale configured, not guessed
    assert "GPA 3.8/4.0" in render(rec)

    # A 10-point GPA keeps its own denominator and shows the conversion too.
    rec["education"][0]["gpa"] = {"raw": 9.14, "scale": 10.0, "normalized_4": 3.66}
    assert "GPA 9.14/10 (≈3.66/4.0)" in render(rec)


def test_fixing_the_gpa_scale_clears_the_stale_warning(tmp_path):
    """Warnings accumulate, which is right for packet findings — but a flag the
    spreadsheet parse re-derives every run must not outlive its cause."""
    import copy
    import yaml as _yaml
    from src.store import Store, db_path_for
    roster = tmp_path / "g.csv"
    roster.write_text("cas_id,first_name,last_name,college_name_0,gpas_by_transcript_gpa_0\n"
                      "1000000001,Ten,Point,B Univ,9.14\n", encoding="utf-8")
    field_map = os.path.join(ROOT, "config", "csv_field_map.yaml")

    run(str(roster), str(tmp_path), field_map)
    rec = Store(db_path_for(str(tmp_path))).get("1000000001")
    assert any("auto-detected" in w for w in rec["warnings"])
    assert "scale not stated" in rec["summary"]

    forced = tmp_path / "forced.yaml"
    cfg = copy.deepcopy(_yaml.safe_load(open(field_map, encoding="utf-8")))
    cfg["gpa"]["gpa_scale"] = 10.0
    forced.write_text(_yaml.safe_dump(cfg), encoding="utf-8")

    run(str(roster), str(tmp_path), str(forced))
    rec = Store(db_path_for(str(tmp_path))).get("1000000001")
    assert not any("auto-detected" in w for w in rec["warnings"])
    assert "GPA 9.14/10 (≈3.66/4.0)" in rec["summary"]


def test_metrics_carry_the_sheets_figure_and_whether_it_was_assumed():
    from src.merge import applicant_metrics, new_student
    u = new_student("G2")
    u["sources"]["application"] = {"education": [
        {"college_name": "X", "gpa": {"raw": 3.8, "scale": 4.0, "normalized_4": 3.8}}]}
    m = applicant_metrics(u)
    assert m["gpa_value"] == "3.8" and m["gpa_assumed"] is False and m["gpa4"] == 3.8

    u["warnings"] = ["GPA scale was auto-detected from the value — verify."]
    assert applicant_metrics(u)["gpa_assumed"] is True


def test_skills_are_not_printed_twice():
    """The résumé's own 'Skills:' line lands in technical preparation, and the
    detected-skills line already says it — printing both put the same list on
    screen twice."""
    from src.documents import build_section
    rec = {
        "cas_id": "CAS-SKILL",
        "personal": {"full_name": "Dup Skills"},
        "_documents": {"resume": build_section("resume", (
            "Education\nState University\nBachelor of Science in Electrical Engineering\n"
            "Skills: Python, Matlab, C++, SPICE\n"
            "Software: OSLO, Lighttools, Code V\n"
            "Research Experience\nSimulated converter controls in Matlab.\n"
        ), "r.txt"), "sop": None, "scores": None, "lors": []},
        "meta": {"warnings": []},
    }
    background = render(rec).split("### Experience")[0]
    assert background.lower().count("skills:") == 1, background
    assert "Software: OSLO" in background          # the additive line survives


def test_undergrad_degree_comes_from_the_packet_only():
    from src.review import undergrad_degree
    resume = ("Education\nMichigan State University\n"
              "Bachelor of Science in Electrical Engineering\nGPA: 3.81\n")
    assert undergrad_degree(_packet_record(resume)) == \
        "Bachelor of Science in Electrical Engineering"
    # No packet -> empty, never inferred from the roster's other columns.
    assert undergrad_degree(_packet_record()) == ""


def test_undergrad_degree_prefers_a_line_that_names_the_field():
    from src.review import undergrad_degree
    # The résumé states the credential but not the subject; the transcript does
    # both, so the transcript wins even though the résumé is checked first.
    got = undergrad_degree(_packet_record(
        resume_text="Education\nIndian School of Mines\nBachelor of Eng.\n",
        scores_text="of this Institute. He/She passed 4-years B.Tech (Electrical Engineering) - with\n"))
    assert got == "B.Tech (Electrical Engineering)"


def test_undergrad_degree_ignores_graduate_only_lines():
    from src.review import undergrad_degree
    assert undergrad_degree(_packet_record(
        resume_text="Education\nMaster of Science in Robotics\nPh.D. in Controls\n")) == ""


def test_undergrad_degree_ignores_transcript_boilerplate():
    """Found by the adversarial pass: a transcript legend line put
    'Undergraduate degree credit' in somebody's degree column."""
    from src.review import undergrad_degree
    for junk in ("NDC = Not for Undergraduate degree credit",
                 "AGC = Approved for Graduate Credit; CBE = Credit by Exam",
                 "This certificate is not a degree"):
        assert undergrad_degree(_packet_record(scores_text=junk)) == "", junk


def test_status_accepts_the_past_tense():
    """Also from the adversarial pass: 'rejected' read as undecided, so a row
    the reviewer had marked silently stayed in the pool."""
    from src.review import normalize_status
    assert normalize_status("rejected") == "reject"
    assert normalize_status("accepted") == "accept"
    assert normalize_status("admitted") == "accept"


def test_status_column_is_read_back_case_and_synonym_insensitively():
    from src.review import normalize_status, status_from_row
    assert status_from_row({"cas_id": "1", "status": "TBD"}) == "tbd"
    assert status_from_row({"cas_id": "1", "Status": " Accept "}) == "accept"
    assert status_from_row({"cas_id": "1", "status": "deny"}) == "reject"
    assert status_from_row({"cas_id": "1"}) == ""
    # A column the reviewer added by hand under the old name still reads back.
    assert status_from_row({"cas_id": "1", "decision": "tbd"}) == "tbd"
    # An unreadable cell is undecided, not a guess.
    assert normalize_status("looks good?") == "" and normalize_status("") == ""


def test_summaries_are_the_reviewed_set():
    """The summaries view and the export carry what you decided in the deck —
    shortlisted first, then rejected. Anything still queued has not been read."""
    from src.run import summaries_markdown
    from src.triage import decided
    entries = [
        {"cas_id": "1", "name": "Zoe", "review_status": "tbd", "summary_text": "z"},
        {"cas_id": "2", "name": "Adam", "review_status": "tbd", "triage": "reject",
         "summary_text": "a"},
        {"cas_id": "3", "name": "Mia", "review_status": "tbd", "triage": "shortlist",
         "summary_text": "m"},
    ]
    assert [e["cas_id"] for e in decided(entries)] == ["3", "2"]   # kept, then cut

    md = summaries_markdown(entries)
    assert "2 of 3 applicants reviewed — 1 shortlisted, 1 rejected" in md
    assert "## Mia  ·  3  —  shortlisted" in md
    assert "## Adam  ·  2  —  rejected" in md
    assert "Zoe" not in md                       # never swiped, never summarized

    assert "Nothing reviewed yet" in summaries_markdown([entries[0]])


def test_review_sheet_round_trip(tmp_path):
    """Roster in -> degree + blank status out -> marked up -> summaries narrow."""
    import openpyxl
    from src.review import write_review_sheet, STATUS_COLUMN, DEGREE_COLUMN
    from src.store import Store, db_path_for

    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    raw = os.path.join(str(tmp_path), "raw")
    os.makedirs(raw, exist_ok=True)
    import shutil
    shutil.copy(CSV, os.path.join(raw, os.path.basename(CSV)))

    store = Store(db_path_for(str(tmp_path)))
    out = write_review_sheet(store.index(), raw, os.path.join(str(tmp_path), "review.xlsx"))
    wb = openpyxl.load_workbook(out)
    ws = wb["Applicants"]
    header = [c.value for c in ws[1]]
    assert STATUS_COLUMN in header and DEGREE_COLUMN in header
    assert "How to review" in wb.sheetnames          # legend, not a fake sample row
    assert ws.max_row == store.count() + 1
    dv = ws.data_validations.dataValidation
    assert dv and "accept,reject,tbd" in dv[0].formula1

    # The reviewer marks one row tbd and the rest rejected, then re-uploads.
    col = header.index(STATUS_COLUMN) + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col).value = "tbd" if r == 2 else "reject"
    marked = os.path.join(raw, "reviewed.xlsx")
    wb.save(marked)
    run(marked, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))

    index = Store(db_path_for(str(tmp_path))).index()
    assert sorted(e["review_status"] for e in index) == ["reject"] * (len(index) - 1) + ["tbd"]
    # The report is the record of what was *reviewed*, and nothing has been
    # swiped yet — marking a spreadsheet row tbd only queues it for the deck.
    report = open(os.path.join(str(tmp_path), "reports", "_summaries.md"),
                  encoding="utf-8").read()
    assert "Nothing reviewed yet" in report and report.count("\n## ") == 0


def test_loose_pdfs_in_a_folder_ingest_like_a_zip(tmp_path):
    """A folder of packets and a ZIP of the same packets take one code path."""
    import pytest
    pytest.importorskip("reportlab")
    from src.ingest_zip import ingest_packet_folder
    from src.store import Store, db_path_for
    folder = tmp_path / "packets"
    folder.mkdir()
    _make_packet(str(folder / "1000419138_Joonha_Jun_Full Application_1_Fall 2020 EE.pdf"),
                 "Joonha Jun",
                 "My research is in power systems and smart grid renewable energy. " * 10,
                 "Education\nState University\nBachelor of Science in Electrical Engineering\n")
    report = ingest_packet_folder(str(folder), str(tmp_path))
    assert report["pdfs_found"] == 1 and report["students_updated"] == 1
    u = Store(db_path_for(str(tmp_path))).get("1000419138")
    assert u["sources"]["sop"] and u["sources"]["resume"]


def test_enrich_folder_writes_the_sheet_next_to_the_roster(tmp_path):
    """The whole of step 1: spreadsheet + packets in a folder -> review sheet
    beside the original, degree filled in from the packet."""
    import pytest
    pytest.importorskip("reportlab")
    import openpyxl
    from src.enrich import enrich_folder, find_roster
    from src.review import STATUS_COLUMN, DEGREE_COLUMN

    folder = tmp_path / "Fall 2027 applicants"
    folder.mkdir()
    # Two applicants; only the first has a packet, so the other's degree cell
    # stays blank and shows up in the coverage counts.
    cas, other = "1000419138", "1000555201"
    (folder / "roster.csv").write_text(
        "cas_id,first_name,last_name,college_name_0,gpas_by_transcript_gpa_0\n"
        f"{cas},Joonha,Jun,State University,3.7\n"
        f"{other},Maria,Lopez,Other University,3.4\n", encoding="utf-8")
    # `_make_packet` draws each section as a single PDF text run, so keep the
    # résumé to one line — embedded newlines come back glued ("...—iBachelor")
    # and that's the fixture's rendering, not what a real packet looks like.
    _make_packet(str(folder / f"{cas}_Joonha_Jun_Full Application_1_Fall 2027 EE.pdf"),
                 "Joonha Jun",
                 "My goal is to research power systems at Texas A&M University. " * 6,
                 "Education: State University. Bachelor of Science in Electrical "
                 "Engineering. Skills: python, matlab.")

    assert find_roster(str(folder)).endswith("roster.csv")
    result = enrich_folder(str(folder), str(tmp_path / "store"))

    out = result["review_sheet"]
    assert out == str(folder / "roster_review.xlsx") and os.path.exists(out)
    assert os.path.exists(folder / "roster.csv"), "the original roster is never modified"

    ws = openpyxl.load_workbook(out)["Applicants"]
    header = [c.value for c in ws[1]]
    assert STATUS_COLUMN in header and DEGREE_COLUMN in header
    degrees = {ws.cell(row=r, column=header.index("cas_id") + 1).value:
               ws.cell(row=r, column=header.index(DEGREE_COLUMN) + 1).value
               for r in range(2, ws.max_row + 1)}
    assert degrees[cas] == "Bachelor of Science in Electrical Engineering"

    # Everyone else had no packet, so their cell is blank — and that's reported
    # as missing data rather than passed off as "no degree".
    cov = result["coverage"]
    assert cov["with_degree"] == 1 and cov["no_packet"] == cov["total"] - 1
    assert "no packet on file" in "\n".join(__import__("src.enrich", fromlist=["describe"])
                                            .describe(result))

    # Re-running picks the roster, not the review sheet it just wrote.
    assert find_roster(str(folder)).endswith("roster.csv")


# ------------------------------------------------- the swipe pass + queues --

def test_deck_card_is_a_trim_of_the_full_summary():
    """The deck card is a trim, not a re-render — one place decides what a
    summary says, this only decides how much of it leads."""
    from src.triage import condense_card
    full = render({
        "cas_id": "CAS-CARD",
        "personal": {"full_name": "Deck Applicant", "email": "d@x.com"},
        "_documents": {"resume": None, "sop": None, "scores": None, "lors": []},
        "meta": {"warnings": ["check this", "and this", "and this too"]},
    })
    card = condense_card(full)
    assert card.splitlines()[0].startswith("Deck Applicant")   # header survives
    assert "### Verify" in card                    # flags matter to the decision
    assert "### On file" not in card               # the header chips say this already
    assert len(card.splitlines()) <= len(full.splitlines())


def test_deck_card_still_caps_a_runaway_section():
    """Loose, not unbounded: a résumé with thirty experience bullets shouldn't
    push everything else off the first screen."""
    from src.triage import condense_card
    text = ("Name · a@b.com\nGPA 4.0/4.0\n\n"
            "### Experience\n" + "\n".join(f"- item {i}" for i in range(30)) + "\n\n"
            "### On file\n- resume 100w")
    card = condense_card(text)
    assert card.count("- item") == 5          # Experience capped
    assert "### On file" not in card          # not a deciding section


def test_swipes_are_separate_from_the_spreadsheet_status():
    from src.triage import apply_swipe
    right = apply_swipe({"review_status": "tbd"}, "shortlist")
    assert right["triage"] == "shortlist" and right["review_status"] == "tbd"

    left = apply_swipe({"review_status": "tbd"}, "reject")
    assert left["triage"] == "reject" and left["review_status"] == "reject"

    back = apply_swipe(dict(left, assigned_to="gratz-paul"), "undo")
    assert "triage" not in back and back["review_status"] == "tbd"
    assert "assigned_to" not in back           # undo clears the assignment too

    import pytest
    with pytest.raises(ValueError):
        apply_swipe({}, "sideways")


def test_deck_holds_tbd_applicants_who_have_not_been_swiped():
    from src.triage import deck, shortlisted
    entries = [
        {"cas_id": "1", "review_status": "tbd"},
        {"cas_id": "2", "review_status": "tbd", "triage": "shortlist"},
        {"cas_id": "3", "review_status": "accept"},
        {"cas_id": "4", "review_status": "reject", "triage": "reject"},
    ]
    assert [e["cas_id"] for e in deck(entries)] == ["1"]
    assert [e["cas_id"] for e in shortlisted(entries)] == ["2"]


def test_professors_load_and_queue_up():
    from src.triage import load_professors, professor_by_id, queues
    profs = load_professors()
    names = [p["name"] for p in profs]
    assert "Paul Gratz" in names and len(profs) >= 4
    assert all(p["url"].startswith("https://") for p in profs if p["url"])
    assert professor_by_id("gratz-paul", profs)["name"] == "Paul Gratz"
    assert professor_by_id("nobody", profs) is None

    entries = [{"cas_id": "1", "triage": "shortlist", "assigned_to": "gratz-paul"},
               {"cas_id": "2", "triage": "shortlist"},                # unassigned
               {"cas_id": "3", "triage": "reject"}]
    filled, unassigned = queues(entries, profs)
    assert [len(p["applicants"]) for p in filled if p["id"] == "gratz-paul"] == [1]
    assert [e["cas_id"] for e in unassigned] == ["2"]


# ------------------------------------------------ accounts, roles, handoff --

def _with_accounts(tmp_path, monkeypatch, password="pw"):
    """Point the app at a throwaway users.yaml and create both roles."""
    from src import users as usersmod
    path = usersmod.USERS_YAML          # conftest points this at tmp_path
    people = usersmod.add_user([], "maya", password, usersmod.ADVISOR)
    people = usersmod.add_user(people, "pgratz", password, usersmod.PROFESSOR,
                               "gratz-paul")
    usersmod.save_users(people, path)
    return path


def test_login_is_off_until_an_account_exists(tmp_path, monkeypatch):
    """A single reviewer shouldn't have to authenticate to themselves — but a
    half-configured login that lets everyone in would be worse than none."""
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    assert not usersmod.load_users()                   # conftest: no accounts
    assert client.get("/").status_code == 200          # -> app is open

    _with_accounts(tmp_path, monkeypatch)
    r = client.get("/")
    assert r.status_code == 302 and "/login" in r.headers["Location"]
    assert client.get("/login").status_code == 200


def test_first_account_is_created_from_the_browser(tmp_path, monkeypatch):
    """Turning logins on shouldn't need the CLI — and the setup page must close
    itself afterwards, or anyone reaching the app could mint an advisor."""
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    path = usersmod.USERS_YAML

    # the banner points at it while the app is wide open
    assert b"Set up logins" in client.get("/").data
    assert client.get("/setup").status_code == 200

    assert b"match" in client.post("/setup", data={
        "username": "maya", "password": "longenough", "confirm": "different"},
        follow_redirects=True).data
    assert b"4 characters" in client.post("/setup", data={
        "username": "maya", "password": "ab", "confirm": "ab"},
        follow_redirects=True).data
    assert usersmod.load_users(path) == []          # nothing written on a bad try

    r = client.post("/setup", data={"username": "maya", "password": "longenough",
                                    "confirm": "longenough"})
    assert r.headers["Location"].endswith("/accounts")
    people = usersmod.load_users(path)
    assert [u["username"] for u in people] == ["maya"] and people[0]["role"] == "advisor"
    assert people[0]["password_hash"] and "longenough" not in people[0]["password_hash"]
    assert client.get("/").status_code == 200       # already signed in

    # closed now: a second visit can't mint another advisor
    r = client.get("/setup")
    assert r.status_code == 302 and "/login" in r.headers["Location"]

    # There is no route for an advisor to set someone else's password — the
    # professor signs up and the advisor approves.
    assert "accounts_add" not in {r.endpoint for r in appmod.app.url_map.iter_rules()}

    client.get("/logout")
    client.post("/signup", data={"username": "pgratz", "password": "longenough",
                                 "confirm": "longenough"})
    client.post("/login", data={"username": "maya", "password": "longenough"})
    client.post("/accounts/approve", data={"username": "pgratz", "role": "professor",
                                           "professor": "gratz-paul"})
    assert {u["username"] for u in usersmod.load_users(path)} == {"maya", "pgratz"}

    # and can't lock themselves out
    r = client.post("/accounts/remove", data={"username": "maya"}, follow_redirects=True)
    assert b"signed in with" in r.data
    assert len(usersmod.load_users(path)) == 2


def test_roles_see_different_apps(tmp_path, monkeypatch):
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)

    # wrong password gets nowhere
    r = client.post("/login", data={"username": "maya", "password": "nope"},
                    follow_redirects=True)
    assert b"don&#39;t match" in r.data or b"don't match" in r.data

    client.post("/login", data={"username": "maya", "password": "pw"})
    assert client.get("/").status_code == 200                 # advisor dashboard
    assert client.get("/step/1").status_code == 200
    assert client.get("/triage").status_code == 200
    client.get("/logout")

    client.post("/login", data={"username": "pgratz", "password": "pw"})
    r = client.get("/")
    assert r.status_code == 302 and r.headers["Location"].endswith("/prof")
    assert client.get("/prof").status_code == 200
    # the advisor's flow is closed to them, loudly rather than silently
    for advisor_only in ("/step/1", "/triage", "/applicants", "/professors"):
        assert client.get(advisor_only).status_code == 403, advisor_only


def test_signup_creates_a_request_not_an_account(tmp_path, monkeypatch):
    """'Create account' must not mean 'grant yourself access'."""
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)

    r = client.post("/signup", data={"username": "stranger", "password": "pw1234",
                                     "confirm": "pw1234"})
    assert r.status_code == 302 and "/login" in r.headers["Location"]

    people = usersmod.load_users()
    row = next(u for u in people if u["username"] == "stranger")
    assert row["role"] == usersmod.PENDING and row["professor"] == ""
    assert row["password_hash"] and "pw1234" not in row["password_hash"]

    # right password, still no way in — and no session is minted
    r = client.post("/login", data={"username": "stranger", "password": "pw1234"},
                    follow_redirects=True)
    assert b"waiting for an advisor" in r.data
    for gated in ("/step/1", "/applicants", "/prof", "/accounts"):
        assert client.get(gated).status_code in (302, 403), gated


def test_pending_signup_cannot_lock_everyone_out(tmp_path, monkeypatch):
    """The trap: an app in no-login mode has no advisor account. If a stranger's
    request counted as 'an account exists', it would flip login on and nobody —
    including the advisor — could get in."""
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    assert client.get("/").status_code == 200          # open, no accounts

    # signup is closed in this state; it hands you to /setup instead
    r = client.post("/signup", data={"username": "x", "password": "pw1234",
                                     "confirm": "pw1234"})
    assert r.status_code == 302 and "/setup" in r.headers["Location"]

    # and even a request written directly doesn't flip the switch
    usersmod.save_users(usersmod.request_account([], "sneaky", "pw1234"))
    assert not usersmod.login_required_globally()
    assert client.get("/").status_code == 200
    assert client.get("/setup").status_code == 200     # advisor can still set up


def test_advisor_approves_and_assigns_the_role(tmp_path, monkeypatch):
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    client.post("/signup", data={"username": "newprof", "password": "pw1234",
                                 "confirm": "pw1234"})

    # a professor can't approve anyone — that's how you'd escalate
    client.post("/login", data={"username": "pgratz", "password": "pw"})
    assert client.post("/accounts/approve",
                       data={"username": "newprof", "role": "advisor"}).status_code == 403
    assert usersmod.load_users()[-1]["role"] == usersmod.PENDING
    client.get("/logout")

    client.post("/login", data={"username": "maya", "password": "pw"})
    page = client.get("/accounts").get_data(as_text=True)
    assert "Waiting for approval" in page and "newprof" in page

    client.post("/accounts/approve", data={"username": "newprof", "role": "professor",
                                           "professor": "duffield-nick"})
    row = next(u for u in usersmod.load_users() if u["username"] == "newprof")
    assert row["role"] == "professor" and row["professor"] == "duffield-nick"

    client.get("/logout")
    client.post("/login", data={"username": "newprof", "password": "pw1234"})
    assert client.get("/prof").status_code == 200
    assert client.get("/applicants").status_code == 403       # still not an advisor


def test_a_professor_login_must_point_at_real_faculty(tmp_path, monkeypatch):
    """An account naming a faculty id that doesn't exist isn't a small
    inconsistency: that person signs in, lands on their deck, and sees an empty
    queue forever with nothing to say why."""
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    client.post("/signup", data={"username": "newprof", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})

    r = client.post("/accounts/approve",
                    data={"username": "newprof", "role": "professor",
                          "professor": "nobody-by-that-name"},
                    follow_redirects=True)
    assert b"No professor with id" in r.data
    assert next(u for u in usersmod.load_users()
                if u["username"] == "newprof")["role"] == usersmod.PENDING

    # the real one works
    client.post("/accounts/approve", data={"username": "newprof",
                                           "role": "professor",
                                           "professor": "fan-zhiwen"})
    row = next(u for u in usersmod.load_users() if u["username"] == "newprof")
    assert row["role"] == "professor" and row["professor"] == "fan-zhiwen"


def test_one_login_per_professor(tmp_path, monkeypatch):
    """Assigning a second account to a professor hands someone else that
    professor's applicants — an access decision wearing a dropdown's clothes."""
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)         # pgratz already holds gratz-paul
    client.post("/signup", data={"username": "impostor", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})

    r = client.post("/accounts/approve",
                    data={"username": "impostor", "role": "professor",
                          "professor": "gratz-paul"}, follow_redirects=True)
    assert b"already signs in as" in r.data
    assert next(u for u in usersmod.load_users()
                if u["username"] == "impostor")["role"] == usersmod.PENDING

    # the taken one isn't even offered
    page = client.get("/accounts").get_data(as_text=True)
    assert 'value="gratz-paul"' not in page
    assert 'value="fan-zhiwen"' in page            # still free

    # freeing it up makes it assignable again
    client.post("/accounts/remove", data={"username": "pgratz"})
    assert 'value="gratz-paul"' in client.get("/accounts").get_data(as_text=True)
    client.post("/accounts/approve", data={"username": "impostor",
                                           "role": "professor",
                                           "professor": "gratz-paul"})
    assert next(u for u in usersmod.load_users()
                if u["username"] == "impostor")["professor"] == "gratz-paul"


def test_when_every_professor_is_taken_the_page_says_so(tmp_path, monkeypatch):
    """The screenshot case: a dropdown that looks full but has nothing in it."""
    from src import triage as triagemod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    triagemod.save_professors([{"id": "gratz-paul", "name": "Paul Gratz"}],
                              triagemod.PROFESSORS_YAML)   # the only one, taken
    client.post("/signup", data={"username": "test", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})

    page = client.get("/accounts").get_data(as_text=True)
    assert "Every professor on the list already has a login" in page
    assert 'value="gratz-paul"' not in page, "the taken one must not be offered"
    assert 'value="__new__"' in page, "but a new colleague is still addable"
    # ...and the role itself must stay choosable, or the advice above it is a
    # lie: the page says "pick someone not on the list yet" while the dropdown
    # refuses to let you be a professor at all.
    assert 'value="professor" disabled' not in page
    assert "disabled" not in page.split("Grant access as")[1].split("</table>")[0]


def test_a_professor_can_be_approved_when_the_faculty_list_is_full(tmp_path, monkeypatch):
    """The end of that road: every existing professor taken, and a new one
    approved anyway by naming them."""
    from src import triage as triagemod
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    triagemod.save_professors([{"id": "gratz-paul", "name": "Paul Gratz"}],
                              triagemod.PROFESSORS_YAML)
    client.post("/signup", data={"username": "newhire", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})

    client.post("/accounts/approve",
                data={"username": "newhire", "role": "professor",
                      "professor": "__new__", "new_professor": "Ada Lovelace"})
    row = next(u for u in usersmod.load_users() if u["username"] == "newhire")
    assert row["role"] == "professor" and row["professor"] == "ada-lovelace"
    assert "Ada Lovelace" in [p["name"] for p in triagemod.load_professors()]


def test_approving_can_create_the_faculty_entry(tmp_path, monkeypatch):
    """A new hire shouldn't mean leaving the page, adding them, and coming back."""
    from src import triage as triagemod
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    client.post("/signup", data={"username": "newhire", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})

    before = {p["id"] for p in triagemod.load_professors()}
    r = client.post("/accounts/approve",
                    data={"username": "newhire", "role": "professor",
                          "professor": "__new__", "new_professor": "Ada Lovelace"},
                    follow_redirects=True)
    assert b"Added Ada Lovelace to the faculty list" in r.data

    added = [p for p in triagemod.load_professors() if p["id"] not in before]
    assert [p["name"] for p in added] == ["Ada Lovelace"]
    row = next(u for u in usersmod.load_users() if u["username"] == "newhire")
    assert row["role"] == "professor" and row["professor"] == added[0]["id"]

    # they land on their own (empty) queue, not the advisor's flow
    client.get("/logout")
    client.post("/login", data={"username": "newhire", "password": "pw1234"})
    assert client.get("/prof").status_code == 200
    assert client.get("/applicants").status_code == 403


def test_creating_faculty_at_approval_rejects_a_duplicate_name(tmp_path, monkeypatch):
    """Two faculty entries for one person would split their queue in half."""
    from src import triage as triagemod
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    client.post("/signup", data={"username": "dupe", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})

    before = len(triagemod.load_professors())
    r = client.post("/accounts/approve",
                    data={"username": "dupe", "role": "professor",
                          "professor": "__new__", "new_professor": "Paul Gratz"},
                    follow_redirects=True)
    assert b"already on the list" in r.data
    assert len(triagemod.load_professors()) == before
    assert next(u for u in usersmod.load_users()
                if u["username"] == "dupe")["role"] == usersmod.PENDING

    # and a blank name doesn't create an unnamed entry
    client.post("/accounts/approve",
                data={"username": "dupe", "role": "professor",
                      "professor": "__new__", "new_professor": "   "},
                follow_redirects=True)
    assert len(triagemod.load_professors()) == before


def test_no_faculty_means_no_professor_account(tmp_path, monkeypatch):
    """What the advisor is told to do when the faculty list is empty."""
    from src import triage as triagemod
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    triagemod.save_professors([], triagemod.PROFESSORS_YAML)   # nobody on file

    client.post("/signup", data={"username": "newprof", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})

    r = client.post("/accounts/approve",
                    data={"username": "newprof", "role": "professor",
                          "professor": "gratz-paul"}, follow_redirects=True)
    assert b"no professors on the faculty list" in r.data.lower()
    assert next(u for u in usersmod.load_users()
                if u["username"] == "newprof")["role"] == usersmod.PENDING

    # the page says so rather than offering an empty dropdown
    page = client.get("/accounts").get_data(as_text=True)
    assert "No faculty on file" in page

    # an advisor account is still fine — it needs no faculty entry
    client.post("/accounts/approve", data={"username": "newprof", "role": "advisor"})
    assert next(u for u in usersmod.load_users()
                if u["username"] == "newprof")["role"] == "advisor"


def test_the_cli_and_the_web_agree_about_faculty(tmp_path, monkeypatch):
    """The check used to live only in the CLI, which is how the web route came
    to accept a made-up id. One implementation, one answer."""
    from src import triage as triagemod
    from src import users as usersmod
    triagemod.save_professors([{"id": "gratz-paul", "name": "Paul Gratz"}],
                              triagemod.PROFESSORS_YAML)
    def refuses(fn, *a):
        try:
            fn(*a)
        except usersmod.UserError:
            return True
        raise AssertionError(f"{a[-1]!r} was accepted as a faculty id")

    for bad in ("", "not-a-real-id"):
        assert refuses(usersmod.add_user, [], "x", "pw1234", usersmod.PROFESSOR, bad)
        assert refuses(usersmod.approve_user,
                       usersmod.request_account([], "y", "pw1234"), "y",
                       usersmod.PROFESSOR, bad)
    assert usersmod.add_user([], "x", "pw1234", usersmod.PROFESSOR, "gratz-paul")


def test_signup_does_not_leak_who_already_has_an_account(tmp_path, monkeypatch):
    """A taken-username error on an open form is a staff directory."""
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    before = len(usersmod.load_users())

    taken = client.post("/signup", data={"username": "maya", "password": "pw1234",
                                         "confirm": "pw1234"}, follow_redirects=True)
    fresh = client.post("/signup", data={"username": "nobody", "password": "pw1234",
                                         "confirm": "pw1234"}, follow_redirects=True)
    assert b"Request sent" in taken.data and b"Request sent" in fresh.data
    assert b"already has an account" not in taken.data

    people = usersmod.load_users()
    assert len(people) == before + 1                    # only the new one landed
    assert next(u for u in people if u["username"] == "maya")["role"] == "advisor"


def test_rejecting_a_request_removes_it(tmp_path, monkeypatch):
    from src import users as usersmod
    appmod, client = _client(tmp_path, monkeypatch)
    _with_accounts(tmp_path, monkeypatch)
    client.post("/signup", data={"username": "spam", "password": "pw1234",
                                 "confirm": "pw1234"})
    client.post("/login", data={"username": "maya", "password": "pw"})
    client.post("/accounts/remove", data={"username": "spam"})
    assert not any(u["username"] == "spam" for u in usersmod.load_users())


def test_approve_rejects_a_role_the_form_did_not_offer(tmp_path, monkeypatch):
    from src import users as usersmod
    users = usersmod.request_account([], "someone", "pw1234")
    for bad in ("", "pending", "admin", "ADVISOR"):
        try:
            usersmod.approve_user(users, "someone", bad)
            raise AssertionError(f"{bad!r} was accepted as a role")
        except usersmod.UserError:
            pass
    # a professor role with no professor named is equally refused
    try:
        usersmod.approve_user(users, "someone", "professor")
        raise AssertionError("professor with no professor id was accepted")
    except usersmod.UserError:
        pass


def test_a_professor_only_rules_on_their_own_queue(tmp_path, monkeypatch):
    appmod, client = _client(tmp_path, monkeypatch)
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    mine, theirs = [e["cas_id"] for e in store.index()][:2]
    store.update(mine, lambda u: {**u, "triage": "shortlist", "assigned_to": "gratz-paul"})
    store.update(theirs, lambda u: {**u, "triage": "shortlist", "assigned_to": "fan-zhiwen"})

    _with_accounts(tmp_path, monkeypatch)
    client.post("/login", data={"username": "pgratz", "password": "pw"})

    page = client.get("/prof").get_data(as_text=True)
    assert page.count('class="tcard"') == 1            # only their own student

    r = client.post("/prof/decide", data={"cas_id": mine, "decision": "accept"})
    assert r.status_code == 200 and r.get_json()["professor_status"] == "accept"

    # someone else's student is refused even with the id in hand
    assert client.post("/prof/decide",
                       data={"cas_id": theirs, "decision": "accept"}).status_code == 403
    assert client.post("/prof/decide",
                       data={"cas_id": mine, "decision": "sideways"}).status_code == 400

    # the professor's call never overwrites the advisor's
    rec = Store(db_path_for(str(tmp_path))).get(mine)
    assert rec["professor_status"] == "accept" and rec["triage"] == "shortlist"
    assert rec["professor_decided_at"]

    client.post("/prof/decide", data={"cas_id": mine, "decision": "undo"})
    assert "professor_status" not in Store(db_path_for(str(tmp_path))).get(mine)


def test_dashboard_shows_each_groups_pipeline(tmp_path, monkeypatch):
    appmod, client = _client(tmp_path, monkeypatch)
    gid = _make_group()
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    ids = [e["cas_id"] for e in store.index()]
    store.update(ids[0], lambda u: {**u, "triage": "shortlist", "group": gid,
                                    "assigned_to": "fan-zhiwen",
                                    "professor_status": "accept",
                                    "final_status": "accept"})
    store.update(ids[1], lambda u: {**u, "triage": "shortlist", "group": gid,
                                    "assigned_to": "fan-zhiwen"})
    store.update(ids[2], lambda u: {**u, "triage": "shortlist"})   # not sent

    page = client.get("/").get_data(as_text=True)
    assert "Systems" in page and "Paul Gratz" in page         # group and its leader
    assert "1 accepted" in page and "1 with reviewers" in page
    assert "Not sent to a group yet" in page


def test_status_report_carries_every_verdict(tmp_path, monkeypatch):
    """One row per applicant with every call on it, and who made it: the
    thing the grad office mails round."""
    import io
    import openpyxl
    from src.groups import load_groups
    from src.status_report import COLUMNS, rows_for, summarize
    from src.store import Store, db_path_for
    appmod, client = _client(tmp_path, monkeypatch)
    gid = _make_group()
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    store = Store(db_path_for(str(tmp_path)))
    ids = [e["cas_id"] for e in store.index()]
    store.update(ids[0], lambda u: {**u, "review_status": "tbd", "triage": "shortlist",
                                    "group": gid, "assigned_to": "fan-zhiwen",
                                    "professor_status": "accept",
                                    "professor_decided_at": "2027-01-05T10:00:00+00:00",
                                    "professor_note": "=strong fit",
                                    "final_status": "reject",
                                    "final_decided_at": "2027-01-07T10:00:00+00:00",
                                    "final_note": "no funding"})
    store.update(ids[1], lambda u: {**u, "review_status": "tbd", "triage": "shortlist",
                                    "group": gid, "assigned_to": "fan-zhiwen"})
    store.update(ids[2], lambda u: {**u, "review_status": "reject", "triage": "reject"})

    r = client.get("/download/status.xlsx")
    assert r.status_code == 200
    assert "applicant_status" in r.headers["Content-Disposition"]

    wb = openpyxl.load_workbook(io.BytesIO(r.data))
    ws = wb["Status"]
    header = [c.value for c in ws[1]]
    assert header == COLUMNS
    assert ws.max_row == len(ids) + 1                    # everyone, not just the decided
    by_id = {row[header.index("cas_id")]: row
             for row in ws.iter_rows(min_row=2, values_only=True)}

    done = by_id[ids[0]]
    assert done[header.index("group")] == "Systems"
    assert done[header.index("group_leader")] == "Paul Gratz"
    assert done[header.index("reviewer")] == "Zhiwen Fan"
    assert done[header.index("first_call")] == "accept"
    assert done[header.index("first_call_date")] == "2027-01-05"
    assert done[header.index("first_call_note")] == "=strong fit"   # text, not a formula
    assert done[header.index("final_call")] == "reject"
    assert done[header.index("final_call_note")] == "no funding"

    # with a reviewer, nothing decided yet
    waiting = by_id[ids[1]]
    assert waiting[header.index("first_call")] == "waiting"
    assert not waiting[header.index("final_call")]
    # never sent anywhere: no group, no calls (an empty cell reads back as None)
    rejected = by_id[ids[2]]
    assert not rejected[header.index("group")] and not rejected[header.index("first_call")]
    assert rejected[header.index("status")] == "reject"

    tallies = {row[0]: row[1:] for row in wb["By group"].iter_rows(min_row=2, values_only=True)}
    # leader, sent, to assign, with reviewers, awaiting final, accepted, rejected
    assert tallies["Systems"] == ("Paul Gratz", 2, 0, 1, 0, 0, 1)

    # the helpers work without openpyxl in the way, and sort by group
    rows = rows_for(store.index(), appmod.load_professors(), load_groups())
    assert rows[0]["group"] == "Systems"
    assert summarize(store.index(), load_groups())[-1]["group"] == "— not sent to a group —"


def _bulk_roster(path, n):
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("cas_id,first_name,last_name,email,college_name_0,"
                 "gpas_by_transcript_gpa_0\n")
        for i in range(n):
            fh.write(f"{1000000000 + i},First{i},Last{i},p{i}@x.edu,"
                     f"State University,{3 + (i % 90) / 100:.2f}\n")
    return str(path)


def test_lists_are_paginated_not_dumped(tmp_path, monkeypatch):
    """Every applicant card on one page was 19 MB of HTML at five thousand —
    fast to serve, unusable to open."""
    appmod, client = _client(tmp_path, monkeypatch)
    run(_bulk_roster(tmp_path / "bulk.csv", 120), str(tmp_path),
        os.path.join(ROOT, "config", "csv_field_map.yaml"))

    page = client.get("/applicants").get_data(as_text=True)
    assert page.count('class="card applicant"') == appmod.PAGE_SIZE
    assert "Page 1 of 3" in page
    assert client.get("/applicants?page=3").get_data(as_text=True).count(
        'class="card applicant"') == 20
    # a page number past the end clamps rather than 500s or shows nothing
    assert 'class="card applicant"' in client.get("/applicants?page=99").get_data(as_text=True)
    assert client.get("/applicants?page=notanumber").status_code == 200


def test_the_deck_streams_in_batches(tmp_path, monkeypatch):
    """The deck holds a batch; the browser asks for more as it works down."""
    appmod, client = _client(tmp_path, monkeypatch)
    run(_bulk_roster(tmp_path / "bulk.csv", 60), str(tmp_path),
        os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    ids = [e["cas_id"] for e in store.index()]
    for sid in ids:
        store.update(sid, lambda u: {**u, "review_status": "tbd"})

    page = client.get("/triage").get_data(as_text=True)
    assert page.count('class="tcard"') == appmod.DECK_BATCH
    # the count is the whole pile, not the batch on screen
    assert ">60</b> to read" in page

    seen = "&".join(f"seen={i}" for i in ids[:appmod.DECK_BATCH])
    more = client.get(f"/triage/cards?{seen}").get_json()
    assert more["html"].count('class="tcard"') == appmod.DECK_BATCH
    assert more["remaining"] == 60 - 2 * appmod.DECK_BATCH
    # and it never hands back someone already swiped past
    assert not any(f'data-cas="{i}"' in more["html"] for i in ids[:appmod.DECK_BATCH])


def test_page_level_counts_do_not_read_every_record(tmp_path, monkeypatch):
    """Coverage and the exports run off the cached index. Reading full records
    (each carrying whole essays) to render a one-line stat is what made the
    status export take eleven seconds."""
    appmod, client = _client(tmp_path, monkeypatch)
    run(_bulk_roster(tmp_path / "bulk.csv", 40), str(tmp_path),
        os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))

    calls = []
    real_all = store.__class__.all
    monkeypatch.setattr(store.__class__, "all",
                        lambda self: (calls.append(1), real_all(self))[1])

    for url in ("/step/1", "/step/2", "/applicants", "/download/status.xlsx",
                "/review-sheet", "/"):
        assert client.get(url).status_code in (200, 302), url
    assert not calls, f"{len(calls)} full-record scans during ordinary page views"


def test_status_report_is_advisor_only(tmp_path, monkeypatch):
    appmod, client = _client(tmp_path, monkeypatch)
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    _with_accounts(tmp_path, monkeypatch)
    client.post("/login", data={"username": "pgratz", "password": "pw"})
    assert client.get("/download/status.xlsx").status_code == 403


def test_missing_professor_file_does_not_break_the_app(tmp_path):
    from src.triage import load_professors
    assert load_professors(str(tmp_path / "nope.yaml")) == []


def test_professor_list_edits_round_trip_through_the_file(tmp_path):
    from src.triage import (add_professor, load_professors, remove_professor,
                            save_professors, update_professor)
    path = str(tmp_path / "professors.yaml")
    profs = add_professor([], "Ada Lovelace", "https://example.edu/ada",
                          "Computing, Analytical Engines")
    save_professors(profs, path)
    loaded = load_professors(path)
    assert loaded == [{"id": "ada-lovelace", "name": "Ada Lovelace",
                       "url": "https://example.edu/ada",
                       "areas": ["Computing", "Analytical Engines"]}]

    # Renaming keeps the id, because assignments point at it.
    renamed = update_professor(loaded, "ada-lovelace", name="Ada B. Lovelace")
    save_professors(renamed, path)
    again = load_professors(path)
    assert again[0]["id"] == "ada-lovelace" and again[0]["name"] == "Ada B. Lovelace"
    assert again[0]["areas"] == ["Computing", "Analytical Engines"]   # untouched

    save_professors(remove_professor(again, "ada-lovelace"), path)
    assert load_professors(path) == []


def test_professor_edits_reject_bad_input():
    import pytest
    from src.triage import (ProfessorError, add_professor, remove_professor,
                            update_professor)
    with pytest.raises(ProfessorError):
        add_professor([], "   ")                                  # no name
    # A javascript: URL would execute from the href we render.
    with pytest.raises(ProfessorError):
        add_professor([], "X", "javascript:alert(1)")
    existing = add_professor([], "Paul Gratz")
    with pytest.raises(ProfessorError):
        add_professor(existing, "paul gratz")                     # duplicate
    with pytest.raises(ProfessorError):
        update_professor(existing, "nobody", name="Y")
    with pytest.raises(ProfessorError):
        remove_professor(existing, "nobody")
    # Two genuinely different people whose names slugify the same still fit.
    two = add_professor(add_professor([], "Li Wei"), "Li-Wei")
    assert [p["id"] for p in two] == ["li-wei", "li-wei-2"]


def test_deleting_a_professor_frees_their_queue(tmp_path, monkeypatch):
    from src import groups as groupsmod
    from src import triage as triagemod
    appmod, client = _client(tmp_path, monkeypatch)
    path = str(tmp_path / "professors.yaml")
    monkeypatch.setattr(triagemod, "PROFESSORS_YAML", path)
    triagemod.save_professors(triagemod.add_professor([], "Paul Gratz"), path)
    gid = _make_group(leader="paul-gratz", members=("paul-gratz",))

    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    cas = store.index()[0]["cas_id"]
    store.update(cas, lambda u: {**u, "review_status": "tbd"})
    client.post("/triage/swipe", data={"cas_id": cas, "direction": "shortlist"})
    client.post("/triage/assign", data={"cas_id": cas, "group": gid})
    client.post(f"/group/{gid}/assign", data={"cas_id": cas, "reviewer": "paul-gratz"})
    assert store.get(cas)["assigned_to"] == "paul-gratz"

    client.post("/professors/delete", data={"id": "paul-gratz"}, follow_redirects=True)
    assert triagemod.load_professors(path) == []
    # Still in the group, back with its leader to reassign, rather than
    # pointing at a professor who no longer exists. And out of the group too.
    rec = store.get(cas)
    assert "assigned_to" not in rec and rec["group"] == gid and rec["triage"] == "shortlist"
    g = groupsmod.load_groups()[0]
    assert g["leader"] == "" and g["members"] == []
    assert "not assigned yet" in client.get("/professors").get_data(as_text=True).lower()


def test_triage_routes_round_trip(tmp_path, monkeypatch):
    """Swipe right, send to a group, see them in that group."""
    appmod, client = _client(tmp_path, monkeypatch)
    gid = _make_group()
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    ids = [e["cas_id"] for e in store.index()]
    for sid in ids:
        store.update(sid, lambda u: {**u, "review_status": "tbd"})

    assert client.get("/triage").status_code == 200
    assert client.get("/triage").get_data(as_text=True).count('class="tcard"') == len(ids)

    r = client.post("/triage/swipe", data={"cas_id": ids[0], "direction": "shortlist"})
    assert r.status_code == 200 and r.get_json()["triage"] == "shortlist"
    client.post("/triage/swipe", data={"cas_id": ids[1], "direction": "reject"})
    # a left swipe takes them out of the tbd pool entirely
    assert client.get("/triage").get_data(as_text=True).count('class="tcard"') == len(ids) - 2

    assert client.post("/triage/swipe",
                       data={"cas_id": ids[2], "direction": "sideways"}).status_code == 400
    assert client.post("/triage/swipe",
                       data={"cas_id": "no-such-id", "direction": "reject"}).status_code == 404

    client.post("/triage/assign", data={"cas_id": ids[0], "group": gid})
    groups_page = client.get("/groups").get_data(as_text=True)
    assert "Systems" in groups_page and "1 sent" in groups_page
    assert client.post("/triage/assign",
                       data={"cas_id": ids[0], "group": "ghost"}).status_code == 400


def test_the_summaries_view_and_download_follow_the_swipe(tmp_path, monkeypatch):
    """Swiping changes what belongs in the report but doesn't re-run the
    pipeline, so the download has to be built fresh rather than served from the
    file the last ingest left on disk."""
    appmod, client = _client(tmp_path, monkeypatch)
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    ids = [e["cas_id"] for e in store.index()]
    for sid in ids:
        store.update(sid, lambda u: {**u, "review_status": "tbd"})

    page = client.get("/summaries").get_data(as_text=True)
    assert "Nothing reviewed yet" in page and page.count('class="summary"') == 0

    client.post("/triage/swipe", data={"cas_id": ids[0], "direction": "shortlist"})
    client.post("/triage/swipe", data={"cas_id": ids[1], "direction": "reject"})

    page = client.get("/summaries").get_data(as_text=True)
    assert page.count('class="summary"') == 2
    assert "shortlisted" in page and "rejected" in page

    md = client.get("/download/summaries").get_data(as_text=True)
    assert "2 of 4 applicants reviewed" in md
    assert md.count("\n## ") == 2                  # the two swiped, nobody else


def test_a_right_swipe_can_name_the_group(tmp_path, monkeypatch):
    """Picking on the card saves a second pass through the shortlist; the
    shortlist can still change it afterwards."""
    appmod, client = _client(tmp_path, monkeypatch)
    gid = _make_group()
    other = _make_group("Power", leader="narayanan-krishna", members=("duffield-nick",))
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    from src.store import Store, db_path_for
    store = Store(db_path_for(str(tmp_path)))
    ids = [e["cas_id"] for e in store.index()]
    for sid in ids:
        store.update(sid, lambda u: {**u, "review_status": "tbd"})

    # the deck offers the picker
    assert 'class="group-pick"' in client.get("/triage").get_data(as_text=True)

    r = client.post("/triage/swipe", data={"cas_id": ids[0], "direction": "shortlist",
                                           "group": gid})
    assert r.get_json()["group"] == gid

    # blank is fine — they land on the shortlist not sent (no field written)
    client.post("/triage/swipe", data={"cas_id": ids[1], "direction": "shortlist",
                                       "group": ""})
    assert store.get(ids[1]).get("group", "") == ""

    # an invented group is refused rather than silently stored
    assert client.post("/triage/swipe", data={"cas_id": ids[2], "direction": "shortlist",
                                              "group": "ghost"}).status_code == 400
    # changing it later, from the shortlist
    client.post("/triage/assign", data={"cas_id": ids[0], "group": other})
    assert store.get(ids[0])["group"] == other
    assert 'name="group"' in client.get("/triage/shortlist").get_data(as_text=True)


def test_review_sheet_keeps_every_original_column(tmp_path):
    import openpyxl
    from src.review import write_review_sheet
    from src.store import Store, db_path_for
    run(CSV, str(tmp_path), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    raw = os.path.join(str(tmp_path), "raw")
    os.makedirs(raw, exist_ok=True)
    import shutil
    shutil.copy(CSV, os.path.join(raw, os.path.basename(CSV)))
    original = list(read_rows(CSV)[0].keys())
    store = Store(db_path_for(str(tmp_path)))
    out = write_review_sheet(store.index(), raw, os.path.join(str(tmp_path), "review.xlsx"))
    header = [c.value for c in openpyxl.load_workbook(out)["Applicants"][1]]
    assert set(original) <= set(header), "the reviewer's own columns must survive"


if __name__ == "__main__":
    import pytest
    raise SystemExit(pytest.main([__file__, "-q"]))


def _make_quick_admit_packet(path: str) -> None:
    """A CAS 'Quick Admit' packet: a resume and a release statement, no SOP.

    The release-statement page carries the spacing CAS sprinkles through its own
    headings in the text layer ("RELEA SE S TATEMENT"), which is exactly what
    defeats the plain string matches in _FORM_TERMS.
    """
    from reportlab.pdfgen import canvas
    from reportlab.lib.pagesizes import LETTER
    c = canvas.Canvas(path, pagesize=LETTER)
    # page 1 — resume
    for i, ln in enumerate(["Prat Applicant",
                            "prat@example.com  +1 469 555 1234",
                            "Dallas, TX",
                            "Education", "Texas A&M University, B.S. Computer Engineering",
                            "Experience: embedded software engineering intern",
                            "Projects: RTL microarchitecture, CMOS adder",
                            "Technical Skills: Verilog, C, Python"]):
        c.setFont("Helvetica", 12); c.drawString(72, 720 - i * 20, ln)
    c.showPage()
    # page 2 — the designations form page that sits between them in a real
    # packet. It matters: it makes the hop into the release statement an
    # ignore->sop transition (cheap) rather than resume->sop (expensive).
    designations = (
        "DESIGNA TIONS|"
        "FALL 2027 QUICK ADMIT M.S. COMPUTER ENGINEERING ( ECEN )|"
        "Organiz ation: Texas A&M Univ ersity      Deliv ery: Onsit e|"
        "Depar tment: Electr ical and Comput er Engineer ing|"
        "Progr am Le vel: Mast ers      Submitt ed Dat e: 08-04- 2026|"
        "Start Term: Fall 2027      Complet ed Dat e: 08-04- 2026|"
        "Campus: College St ation      Verified Dat e: -"
    ).split("|")
    for i, ln in enumerate(designations):
        c.setFont("Helvetica", 11); c.drawString(60, 720 - i * 16, ln)
    c.showPage()
    # page 3 — release statement, first-person boilerplate, garbled heading
    release = (
        "RELEA SE S TATEMENT|"
        "By accepting these t erms, I cer tify, as r equir ed in the applic ation, that I|"
        "have read, underst and, and agr ee to all policies f ound within the Help Cent er,|"
        "including the pr ovisions that place r esponsibilit y for monit oring and ensur ing|"
        "the pr ogress of m y applic ation pr ocess with me. I cer tify that all of the|"
        "information and st atements I ha ve pro vided as p art of this applic ation pr ocess|"
        "are current, accur ate, and complet e to the best of m y knowledge. I underst and|"
        "that withholding inf ormation r equest ed as p art of the applic ation pr ocess ma y|"
        "be grounds f or denial of admission b y any p articip ating pr ogram. I agr ee that|"
        "EngineeringC AS ma y release all applic ation inf ormation I ha ve submitt ed to the|"
        "graduat e progr ams I designat e, and I pr ovide consent f or educ ational r esearch."
    ).split("|")
    for i, ln in enumerate(release):
        c.setFont("Helvetica", 11); c.drawString(60, 720 - i * 16, ln)
    c.showPage()
    c.save()


def test_release_statement_is_not_mistaken_for_an_sop(tmp_path):
    """A CAS release statement is first-person boilerplate, not a statement of
    purpose. On a Quick Admit packet (no SOP at all) it used to become the whole
    sop section and suppress the missing-SOP flag, so the reviewer saw legal
    boilerplate presented as the applicant's statement."""
    import pytest
    pytest.importorskip("reportlab")
    from src.packet import segment
    pdf = tmp_path / "quick_admit.pdf"
    _make_quick_admit_packet(str(pdf))
    seg = segment(str(pdf))
    assert "resume" in seg["sections"]
    assert "sop" not in seg["sections"], seg["sections"].get("sop", {}).get("text", "")[:200]
    assert any("statement-of-purpose" in f for f in seg["flags"]), seg["flags"]
