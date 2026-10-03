"""The prescreen spreadsheets the grad office asked for.

Three files, because MS Distance and PhD Distance applications were asked to
live in their own files:

* `prescreen_summary.xlsx`: everyone else (Quick Admit and the other on-campus
  programs), one row each with the fields the office listed, the rule outcome
  and the full written summary;
* `prescreen_ms_distance.xlsx` and `prescreen_phd_distance.xlsx`: the distance
  applicants, with a column per requirement so "which requirement isn't met"
  is readable straight across the row.

Every row starts with a clickable link to that applicant's page in the app.
The link is only as reachable as the app: on a laptop it opens on that laptop;
on the hosted app it opens for anyone with a login. Set GRADAPP_PUBLIC_URL when
the address people use differs from the one the request came in on.

Built from index entries (the cached per-applicant read model), never from
full records, so a few thousand rows is a lookup, not a re-parse. Values, not
formulas: applicant-supplied text that happens to start with "=" is written as
text so a résumé line can't become a live formula in someone's Excel.
"""
from __future__ import annotations

import io
import os
import re
import zipfile
from datetime import datetime, timezone
from typing import Any, Dict, Iterable, List, Sequence, Tuple
from urllib.parse import quote

from .screening import (ADMIT, MS_DISTANCE, PHD_DISTANCE, REVIEW, ON_CAMPUS,
                        OTHER_DISTANCE, NO_PROGRAM, QUICK_ADMIT, courses_source,
                        load_rules)

FILES = {
    "summary": "prescreen_summary.xlsx",
    "ms-distance": "prescreen_ms_distance.xlsx",
    "phd-distance": "prescreen_phd_distance.xlsx",
}

LINK = "Open"
_TRACK_ORDER = {QUICK_ADMIT: 0, ON_CAMPUS: 1, OTHER_DISTANCE: 2, NO_PROGRAM: 3}
_OUTCOME_ORDER = {ADMIT: 0, REVIEW: 1, "": 2}

COMMON = ["CAS ID", "Name", "Email", "Program"]
FIELDS = ["Domestic / International", "Schools & GPA", "Degree type", "Major",
          "GRE", "English proficiency", "Awards"]

SUMMARY_COLUMNS = ([LINK] + COMMON + ["Track"] + FIELDS
                   + ["Prescreen", "Why", "Check by hand"])
MS_COLUMNS = ([LINK] + COMMON + ["Thesis / non-thesis"] + FIELDS
              + ["ECE background", "GPA used for the rule", "Courses listed",
                 "Approved 400-level", "Approved 600/700-level", "Courses counted",
                 "Research advisor proof", "Employer letter / statement",
                 "Prescreen", "Requirements not met", "Check by hand"])
PHD_COLUMNS = ([LINK] + COMMON + FIELDS
               + ["Holds an MS?", "Courses required", "Courses listed",
                  "Approved 400-level", "Approved 600/700-level", "Courses counted",
                  "Employer letter", "Prescreen", "Requirements not met",
                  "Check by hand"])

# Everything wraps except the link. The full written summary is pages long, so
# it gets its own "Summaries" tab rather than a column that would make every
# row on the main tab screen-high.
_NO_WRAP = {LINK}
_WIDTHS = {LINK: 8, "CAS ID": 13, "Name": 22, "Email": 26, "Program": 38,
           "Track": 18, "Thesis / non-thesis": 14, "Domestic / International": 14,
           "Schools & GPA": 44, "Degree type": 26, "Major": 24, "GRE": 18,
           "English proficiency": 20, "Awards": 40, "Prescreen": 14, "Why": 48,
           "Check by hand": 44, "Summary": 80, "Requirements not met": 48,
           "ECE background": 26, "GPA used for the rule": 24, "Courses listed": 36,
           "Approved 400-level": 16, "Approved 600/700-level": 30,
           "Courses counted": 12, "Research advisor proof": 30,
           "Employer letter / statement": 30, "Employer letter": 30,
           "Holds an MS?": 26, "Courses required": 12}


def link_for(base_url: str, cas_id: str) -> str:
    return f"{(base_url or '').rstrip('/')}/student/{quote(str(cas_id), safe='')}"


def _lines(items: Iterable[str]) -> str:
    return "\n".join(str(i) for i in items if i)


def _plain_summary(text: str) -> str:
    """The written summary as plain lines for a cell: Markdown headings become
    "Heading:" so the text reads the same in Excel as on screen."""
    out = []
    for line in (text or "").splitlines():
        m = re.match(r"^#{1,6}\s+(.*)$", line)
        out.append(f"{m.group(1)}:" if m else line)
    return "\n".join(out).strip()[:32000]       # Excel's per-cell limit is 32,767


def _common(entry: Dict[str, Any], base_url: str) -> Dict[str, Any]:
    prof = entry.get("profile") or {}
    scr = entry.get("screening") or {}
    return {
        LINK: link_for(base_url, entry.get("cas_id", "")),
        "CAS ID": entry.get("cas_id", ""),
        "Name": entry.get("name", ""),
        "Email": entry.get("email", ""),
        "Program": scr.get("program") or "; ".join(entry.get("programs") or []),
        "Domestic / International": prof.get("residency", ""),
        "Schools & GPA": _lines(prof.get("schools") or []),
        "Degree type": prof.get("degree_type", ""),
        "Major": prof.get("major", ""),
        "GRE": prof.get("gre", ""),
        "English proficiency": prof.get("english", ""),
        "Awards": _lines(prof.get("awards") or []),
        "Prescreen": scr.get("outcome", ""),
        "Check by hand": _lines(scr.get("verify") or []),
        "Summary": _plain_summary(entry.get("summary_text", "")),
    }


def _course_cols(scr: Dict[str, Any]) -> Dict[str, Any]:
    c = (scr.get("checks") or {}).get("courses") or {}
    listed = c.get("listed") or []
    src = c.get("source") or ""
    return {
        "Courses listed": (", ".join(listed) + (f"\n(from {src})" if src else ""))
                          if listed else "None found",
        "Approved 400-level": ", ".join(c.get("level_400") or []),
        "Approved 600/700-level": ", ".join(c.get("level_600_700") or []),
        "Courses counted": f"{c.get('counted', 0)} of {c.get('required', '')}",
    }


def _letter_col(found: Dict[str, Any]) -> str:
    if not found:
        return ""
    status = found.get("status")
    opt = found.get("option", "")
    prefix = f"Option {opt}: " if opt and not opt.startswith("PhD") else ""
    if status == "exact":
        return f"{prefix}exact wording ({found.get('where')})"
    if status == "close":
        return f"{prefix}close, not exact ({found.get('where')}, " \
               f"{int(float(found.get('coverage', 0)) * 100)}%)"
    return "Not found"


def summary_rows(entries: Sequence[Dict[str, Any]], base_url: str) -> List[Dict[str, Any]]:
    rows = []
    for e in entries:
        scr = e.get("screening") or {}
        row = _common(e, base_url)
        row["Track"] = scr.get("track_label", "")
        row["Why"] = _lines(scr.get("reasons") or [])
        row["_sort"] = (_TRACK_ORDER.get(scr.get("track"), 9),
                        _OUTCOME_ORDER.get(scr.get("outcome", ""), 9),
                        str(e.get("name", "")).lower())
        rows.append(row)
    return sorted(rows, key=lambda r: r["_sort"])


def ms_rows(entries: Sequence[Dict[str, Any]], base_url: str) -> List[Dict[str, Any]]:
    rows = []
    for e in entries:
        scr = e.get("screening") or {}
        checks = scr.get("checks") or {}
        row = _common(e, base_url)
        thesis = scr.get("thesis")
        row["Thesis / non-thesis"] = ("Thesis" if thesis else "Non-thesis"
                                      if thesis is False else "Not stated")
        ece = checks.get("ece") or {}
        row["ECE background"] = {"ece": "Yes", "other": "No", "unknown": "Not found"}.get(
            ece.get("status"), "") + (f": {'; '.join(ece.get('majors') or [])}"
                                      if ece.get("majors") else "")
        gpa = checks.get("gpa") or {}
        row["GPA used for the rule"] = gpa.get("text") or "None on file"
        row.update(_course_cols(scr))
        if thesis:
            adv = checks.get("advisor") or {}
            row["Research advisor proof"] = (f"Found ({adv.get('where')})"
                                             if adv.get("status") == "found" else "Not found")
            row["Employer letter / statement"] = _letter_col(checks.get("employer_letter") or {})
        else:
            row["Research advisor proof"] = "Not required" if thesis is False else ""
            row["Employer letter / statement"] = "Not required" if thesis is False else ""
        row["Requirements not met"] = _lines(scr.get("reasons") or [])
        row["_sort"] = ({False: 0, True: 1}.get(thesis, 2),
                        _OUTCOME_ORDER.get(scr.get("outcome", ""), 9),
                        str(e.get("name", "")).lower())
        rows.append(row)
    return sorted(rows, key=lambda r: r["_sort"])


def phd_rows(entries: Sequence[Dict[str, Any]], base_url: str) -> List[Dict[str, Any]]:
    rows = []
    for e in entries:
        scr = e.get("screening") or {}
        checks = scr.get("checks") or {}
        row = _common(e, base_url)
        ms = checks.get("masters") or {}
        has = ms.get("has_ms")
        row["Holds an MS?"] = ({True: "Yes", False: "No (Bachelor's)"}.get(has, "Couldn't tell")
                               + (f": {ms['evidence']}" if ms.get("evidence") else ""))
        row.update(_course_cols(scr))
        c = checks.get("courses") or {}
        row["Courses required"] = c.get("required", "")
        row["Employer letter"] = _letter_col(checks.get("employer_letter") or {})
        row["Requirements not met"] = _lines(scr.get("reasons") or [])
        row["_sort"] = (_OUTCOME_ORDER.get(scr.get("outcome", ""), 9),
                        str(e.get("name", "")).lower())
        rows.append(row)
    return sorted(rows, key=lambda r: r["_sort"])


def split(entries: Iterable[Dict[str, Any]]) -> Tuple[List[Dict], List[Dict], List[Dict]]:
    """(everyone else, MS Distance, PhD Distance)."""
    main, ms, phd = [], [], []
    for e in entries:
        track = (e.get("screening") or {}).get("track")
        (ms if track == MS_DISTANCE else phd if track == PHD_DISTANCE else main).append(e)
    return main, ms, phd


def rules_text(kind: str) -> List[Tuple[str, str]]:
    """The rules as applied, in words, for the workbook's Rules sheet."""
    r = load_rules()
    c = r.get("courses") or {}
    src = courses_source()
    gpa = r.get("gpa_threshold", 3.25)
    majors = ", ".join(r.get("ece_majors") or [])
    course_rule = (f"Courses must be on the approved distance list ({src.get('url')}, "
                   f"copied {src.get('retrieved')}). At most {c.get('max_400_level', 2)} "
                   "400-level courses; the rest 600/700-level.")
    lines = [("How to read this", "Prescreen is a flag, not a decision. 'Admit' means every "
              "rule was met on the data the app could read; 'Further review' lists what wasn't "
              "met or couldn't be confirmed. 'Check by hand' lists what a person should still "
              "eyeball even on an Admit (letterhead, signatures, the advisor being TAMU faculty)."),
             ("Link", "The first column opens the applicant's page in the app (sign in first).")]
    if kind == "summary":
        lines += [("Quick Admit (on campus)",
                   f"Admit: from the ECE department and GPA {gpa} or higher. Further review: "
                   f"ECE with GPA below {gpa}, or any other department regardless of GPA."),
                  ("Other on-campus programs", "Summarized only; no rule was given for them.")]
    if kind == "ms-distance":
        lines += [("Non-thesis", f"Admit: ECE, GPA {gpa}+, and {c.get('ms_required', 10)} "
                   f"approved courses. {course_rule}"),
                  ("Thesis", "Everything for non-thesis, plus a document from a faculty member "
                   "confirming they'll advise, plus the employer letter (Option A) or the "
                   "unemployment statement (Option B) with the exact required wording.")]
    if kind == "phd-distance":
        lines += [("PhD Distance", f"Admit: {c.get('phd_required_with_ms', 6)} approved courses "
                   f"with an MS, {c.get('phd_required_bachelors', 14)} with a Bachelor's only, "
                   f"plus the PhD employer letter with the exact required wording. {course_rule}")]
    lines += [("GPA", "CAS's cumulative undergraduate GPA when CAS calculated one; otherwise the "
               "4.0-scale GPA(s) on file. A GPA on another scale (e.g. 9.1/10) is never "
               "converted to make the call; it goes to further review."),
              ("ECE department", f"Major / field of study matches: {majors}."),
              ("Letters", "Wording is compared word for word, in order, ignoring case, spacing "
               "and punctuation. 'Close' means most of it matches; compare it by hand."),
              ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))]
    return lines


def write_workbook(path: str, title: str, columns: Sequence[str],
                   rows: Sequence[Dict[str, Any]], kind: str) -> str:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    font, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    link_font = Font(name="Arial", size=10, color="0563C1", underline="single")
    head_fill = PatternFill("solid", fgColor="EEF1F6")
    tint = {ADMIT: PatternFill("solid", fgColor="E6F6EF"),
            REVIEW: PatternFill("solid", fgColor="FDF3DD")}
    wrap = Alignment(wrap_text=True, vertical="top")
    top = Alignment(vertical="top")

    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(list(columns))
    for cell in ws[1]:
        cell.font, cell.fill = bold, head_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    outcome_col = columns.index("Prescreen") if "Prescreen" in columns else None
    for r_i, row in enumerate(rows, start=2):
        for c_i, col in enumerate(columns, start=1):
            value = row.get(col, "")
            cell = ws.cell(row=r_i, column=c_i)
            if col == LINK:
                cell.value = "Open"
                cell.hyperlink = value
                cell.font = link_font
                cell.alignment = top
                continue
            cell.value = value if value != "" else None
            if cell.data_type == "f":       # never a live formula from applicant text
                cell.data_type = "s"
            cell.font = font
            cell.alignment = top if col in _NO_WRAP else wrap
        if outcome_col is not None:
            shade = tint.get(row.get("Prescreen", ""))
            if shade:
                ws.cell(row=r_i, column=outcome_col + 1).fill = shade
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = _WIDTHS.get(col, 16)
    ws.freeze_panes = ws.cell(row=2, column=4)       # keep link, id and name in view
    ws.auto_filter.ref = ws.dimensions
    # Printable as-is: landscape, every column on one page width, header row
    # repeated on each page.
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"

    # The written summaries, one applicant per row, same order and same link.
    summ = wb.create_sheet("Summaries")
    summ.append([LINK, "CAS ID", "Name", "Prescreen", "Summary"])
    for cell in summ[1]:
        cell.font, cell.fill = bold, head_fill
    for r_i, row in enumerate(rows, start=2):
        link = summ.cell(row=r_i, column=1, value="Open")
        link.hyperlink, link.font, link.alignment = row.get(LINK, ""), link_font, top
        for c_i, col in enumerate(["CAS ID", "Name", "Prescreen", "Summary"], start=2):
            cell = summ.cell(row=r_i, column=c_i, value=row.get(col) or None)
            if cell.data_type == "f":
                cell.data_type = "s"
            cell.font = font
            cell.alignment = wrap if col == "Summary" else top
    for letter, width in zip("ABCDE", (8, 13, 22, 14, 120)):
        summ.column_dimensions[letter].width = width
    summ.freeze_panes = "B2"
    summ.page_setup.orientation = "landscape"
    summ.sheet_properties.pageSetUpPr.fitToPage = True
    summ.page_setup.fitToWidth = 1
    summ.page_setup.fitToHeight = 0

    guide = wb.create_sheet("Rules")
    for label, text in rules_text(kind):
        guide.append([label, text])
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 110
    for row in guide.iter_rows():
        row[0].font = bold
        row[0].alignment = top
        row[1].font = font
        row[1].alignment = wrap

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path


def write_all(entries: Iterable[Dict[str, Any]], base_url: str, outdir: str) -> Dict[str, str]:
    """Write all three workbooks into `outdir`. Returns {kind: path}."""
    main, ms, phd = split(entries)
    return {
        "summary": write_workbook(os.path.join(outdir, FILES["summary"]), "Applicants",
                                  SUMMARY_COLUMNS, summary_rows(main, base_url), "summary"),
        "ms-distance": write_workbook(os.path.join(outdir, FILES["ms-distance"]), "MS Distance",
                                      MS_COLUMNS, ms_rows(ms, base_url), "ms-distance"),
        "phd-distance": write_workbook(os.path.join(outdir, FILES["phd-distance"]),
                                       "PhD Distance", PHD_COLUMNS, phd_rows(phd, base_url),
                                       "phd-distance"),
    }


def zip_bytes(paths: Dict[str, str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in paths.values():
            zf.write(path, os.path.basename(path))
    return buf.getvalue()


def counts(entries: Iterable[Dict[str, Any]]) -> Dict[str, Dict[str, int]]:
    """{track: {"total", "admit", "review"}} for the page header."""
    out: Dict[str, Dict[str, int]] = {}
    for e in entries:
        scr = e.get("screening") or {}
        t = out.setdefault(scr.get("track", NO_PROGRAM), {"total": 0, "admit": 0, "review": 0})
        t["total"] += 1
        if scr.get("outcome") == ADMIT:
            t["admit"] += 1
        elif scr.get("outcome") == REVIEW:
            t["review"] += 1
    return out
