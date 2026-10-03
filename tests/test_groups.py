"""Review groups: the grad office sends applicants to a group, the leader
hands them to members, a member makes the first call, the leader makes the
final call. Pure rules first, then the whole flow through the app with real
logins so the permission boundaries are tested, not assumed."""
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from src import groups as G
from src.triage import load_professors

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CSV = os.path.join(ROOT, "tests", "fixtures", "sample_applications.csv")
FIELD_MAP = os.path.join(ROOT, "config", "csv_field_map.yaml")

SYSTEMS = {"id": "systems", "name": "Systems", "leaders": ["gratz-paul"],
           "members": ["fan-zhiwen", "duffield-nick"], "areas": ["Networking"]}


# ------------------------------------------------------------------- rules

def test_group_edits_are_vetted_against_the_faculty_list():
    profs = load_professors()
    groups = G.add_group([], "Systems", leaders=["gratz-paul"],
                         members=["fan-zhiwen", "fan-zhiwen"], professors=profs)
    assert groups[0]["id"] == "systems" and groups[0]["members"] == ["fan-zhiwen"]
    with pytest.raises(G.GroupError):
        G.add_group(groups, "systems", professors=profs)          # duplicate name
    with pytest.raises(G.GroupError):
        G.add_group(groups, "Power", leaders=["nobody"], professors=profs)
    with pytest.raises(G.GroupError):
        G.update_group(groups, "systems", members=["ghost"], professors=profs)
    renamed = G.update_group(groups, "systems", name="Comp Eng", professors=profs)
    assert renamed[0]["id"] == "systems"                        # id never changes
    assert G.drop_professor(renamed, "gratz-paul")[0]["leaders"] == []
    assert G.remove_group(renamed, "systems") == []


def test_round_trip_through_the_file(tmp_path):
    path = str(tmp_path / "g.yaml")
    G.save_groups([SYSTEMS], path)
    assert G.load_groups(path) == [SYSTEMS]
    assert G.load_groups(str(tmp_path / "missing.yaml")) == []


def test_the_leader_can_review_too():
    assert G.reviewers(SYSTEMS) == ["fan-zhiwen", "duffield-nick", "gratz-paul"]


def test_moving_groups_starts_the_review_over():
    u = {"group": "systems", "assigned_to": "fan-zhiwen", "professor_status": "accept",
         "final_status": "accept"}
    G.send_to_group(u, "power")
    assert u == {"group": "power"}
    # an applicant handed straight to a professor before groups existed keeps
    # that reviewer when first sent to a group
    legacy = {"assigned_to": "fan-zhiwen", "professor_status": "accept"}
    G.send_to_group(legacy, "systems")
    assert legacy["assigned_to"] == "fan-zhiwen" and legacy["professor_status"] == "accept"
    G.send_to_group(legacy, "")
    assert legacy == {}


def test_reviewer_rules():
    u = {"group": "systems"}
    with pytest.raises(G.GroupError):
        G.assign_reviewer(u, "narayanan-krishna", SYSTEMS)       # not in the group
    G.assign_reviewer(u, "fan-zhiwen", SYSTEMS)
    G.apply_first_call(u, "accept", "  great   fit ")
    assert u["professor_note"] == "great fit"
    G.assign_reviewer(u, "duffield-nick", SYSTEMS)               # new reviewer, fresh call
    assert "professor_status" not in u and "professor_note" not in u
    with pytest.raises(G.GroupError):
        G.assign_reviewer({"group": "power"}, "fan-zhiwen", SYSTEMS)


def test_final_call_comes_after_the_first():
    u = {"group": "systems"}
    with pytest.raises(G.GroupError, match="Assign a reviewer"):
        G.apply_final_call(u, "accept", leader_id="gratz-paul", by="x", group=SYSTEMS)
    G.assign_reviewer(u, "fan-zhiwen", SYSTEMS)
    with pytest.raises(G.GroupError, match="waiting on the reviewer"):
        G.apply_final_call(u, "accept", leader_id="gratz-paul", by="x", group=SYSTEMS)
    G.apply_first_call(u, "accept")
    G.apply_final_call(u, "reject", leader_id="gratz-paul", by="pgratz",
                       note="no funding", group=SYSTEMS)
    assert u["final_status"] == "reject" and u["professor_status"] == "accept"
    assert u["final_by"] == "pgratz" and u["final_note"] == "no funding"
    # the reviewer can't change their call under a final one, nor can the
    # reviewer be swapped, until the leader undoes it
    with pytest.raises(G.GroupError):
        G.apply_first_call(u, "reject")
    with pytest.raises(G.GroupError):
        G.assign_reviewer(u, "duffield-nick", SYSTEMS)
    G.apply_final_call(u, "undo", leader_id="gratz-paul", by="pgratz", group=SYSTEMS)
    assert "final_status" not in u and u["professor_status"] == "accept"


def test_a_leader_reviewing_their_own_applicant_calls_it_once():
    u = {"group": "systems", "assigned_to": "gratz-paul"}
    G.apply_final_call(u, "accept", leader_id="gratz-paul", by="pgratz", group=SYSTEMS)
    assert u["professor_status"] == u["final_status"] == "accept"


def test_stages():
    assert G.stage({}) == G.NOT_SENT
    assert G.stage({"group": "s"}) == G.TO_DISTRIBUTE
    assert G.stage({"group": "s", "assigned_to": "x"}) == G.WITH_REVIEWER
    assert G.stage({"group": "s", "assigned_to": "x", "professor_status": "accept"}) \
        == G.READY_FOR_FINAL
    assert G.stage({"group": "s", "final_status": "reject"}) == G.DECIDED


def test_split_evenly_levels_open_load():
    entries = [{"cas_id": str(i), "name": f"N{i}", "group": "systems"} for i in range(5)]
    # duffield already has one open review, fan has a finished one
    entries += [{"cas_id": "x", "name": "X", "group": "systems", "assigned_to": "duffield-nick"},
                {"cas_id": "y", "name": "Y", "group": "systems", "assigned_to": "fan-zhiwen",
                 "professor_status": "accept", "final_status": "accept"}]
    plan = G.split_evenly(entries, SYSTEMS)
    got = [who for _cas, who in plan]
    assert len(plan) == 5 and "gratz-paul" not in got           # leader isn't a member
    assert got.count("fan-zhiwen") == 3 and got.count("duffield-nick") == 2


def test_split_breaks_ties_by_area():
    profs = [{"id": "fan-zhiwen", "areas": ["Computer Vision"]},
             {"id": "duffield-nick", "areas": ["Networking"]}]
    entries = [{"cas_id": "1", "name": "A", "group": "systems",
                "metrics": {"areas": ["Networking"]}}]
    assert G.split_evenly(entries, SYSTEMS, profs) == [("1", "duffield-nick")]


def test_suggest_group_from_areas():
    power = {"id": "power", "name": "Power", "areas": ["Power Systems"]}
    entry = {"metrics": {"areas": ["power systems", "Controls"]}}
    assert G.suggest_group(entry, [SYSTEMS, power]) == "power"
    assert G.suggest_group({"metrics": {"areas": ["Photonics"]}}, [SYSTEMS, power]) == ""


# ---------------------------------------------------- the flow, with logins

@pytest.fixture
def office(tmp_path, monkeypatch):
    """Roster loaded and four logins: the grad office, a Group leader (who
    will lead Systems), a member, and a professor from outside the group."""
    from src import app as appmod
    from src import users as usersmod
    from src.run import run
    from src.store import Store, db_path_for
    monkeypatch.setattr(appmod, "DATA", str(tmp_path))
    os.makedirs(tmp_path / "raw", exist_ok=True)
    appmod.app.config.update(TESTING=True)
    run(CSV, str(tmp_path), FIELD_MAP)
    store = Store(db_path_for(str(tmp_path)))
    ids = [e["cas_id"] for e in store.index()]
    for cas in ids:
        store.update(cas, lambda u: {**u, "review_status": "tbd", "triage": "shortlist"})
    people = usersmod.add_user([], "office", "pw", usersmod.ADVISOR)
    people = usersmod.add_user(people, "pgratz", "pw", usersmod.LEADER, "gratz-paul")
    people = usersmod.add_user(people, "pfan", "pw", usersmod.PROFESSOR, "fan-zhiwen")
    people = usersmod.add_user(people, "pkrishna", "pw", usersmod.PROFESSOR,
                               "narayanan-krishna")
    usersmod.save_users(people, usersmod.USERS_YAML)
    return appmod, appmod.app.test_client(), store, ids


def _as(client, who):
    client.get("/logout")
    r = client.post("/login", data={"username": who, "password": "pw"})
    assert r.status_code == 302, who


def test_the_whole_flow(office):
    appmod, client, store, ids = office
    cas, other = ids[0], ids[1]
    fetch = {"X-Requested-With": "fetch"}

    # grad office sets up the group and sends an applicant to it
    _as(client, "office")
    client.post("/groups/add", data={"name": "Systems", "leaders": ["gratz-paul"],
                                     "members": ["fan-zhiwen", "duffield-nick"]})
    assert [g["id"] for g in G.load_groups()] == ["systems"]
    assert client.post("/triage/assign", data={"cas_id": cas, "group": "systems"},
                       headers=fetch).get_json()["group"] == "systems"
    # the office can look at the group, but the hand-out and final call are the leader's
    assert client.get("/group/systems").status_code == 200
    assert client.post("/group/systems/assign",
                       data={"cas_id": cas, "reviewer": "fan-zhiwen"}).status_code == 403

    # before anyone hands it out, members and outsiders can't open the file
    _as(client, "pfan")
    assert client.get(f"/student/{cas}").status_code == 403
    assert client.get("/group/systems").status_code == 403

    # the leader sees the group from their own page and hands the applicant out
    _as(client, "pgratz")
    home = client.get("/prof").get_data(as_text=True)
    assert "Groups you lead" in home and "1 to assign" in home
    assert client.get(f"/student/{cas}").status_code == 200       # leader can read it
    assert client.get(f"/student/{other}").status_code == 403     # not in their group
    r = client.post("/group/systems/assign",
                    data={"cas_id": cas, "reviewer": "narayanan-krishna"}, headers=fetch)
    assert r.status_code == 409                                   # not a member
    r = client.post("/group/systems/final", data={"cas_id": cas, "decision": "accept"},
                    headers=fetch)
    assert r.status_code == 409                                   # nobody's reviewed it
    client.post("/group/systems/assign", data={"cas_id": cas, "reviewer": "fan-zhiwen"})
    assert store.get(cas)["assigned_to"] == "fan-zhiwen"

    # the member makes the first call, with a note for the leader
    _as(client, "pfan")
    assert client.get("/prof").get_data(as_text=True).count('class="tcard"') == 1
    assert client.get(f"/student/{cas}").status_code == 200
    r = client.post("/prof/decide", data={"cas_id": cas, "decision": "accept",
                                          "note": "strong systems background"})
    assert r.get_json()["professor_status"] == "accept"

    # someone outside the group still can't open it
    _as(client, "pkrishna")
    assert client.get(f"/student/{cas}").status_code == 403
    assert client.post("/prof/decide",
                       data={"cas_id": cas, "decision": "reject"}).status_code == 403

    # the leader sees the member's call and note, and makes the final call
    _as(client, "pgratz")
    page = client.get("/group/systems").get_data(as_text=True)
    assert "strong systems background" in page and "Ready for your final call" in page
    client.post("/group/systems/final", data={"cas_id": cas, "decision": "reject",
                                              "note": "no funding this cycle"})
    rec = store.get(cas)
    assert rec["final_status"] == "reject" and rec["final_by"] == "pgratz"
    assert rec["professor_status"] == "accept"                    # both calls kept

    # the member can't quietly change their call under the final one
    _as(client, "pfan")
    assert client.post("/prof/decide",
                       data={"cas_id": cas, "decision": "undo"}).status_code == 409

    # the grad office sees both calls, the notes, and who decided
    _as(client, "office")
    assert "no funding this cycle" in client.get(f"/student/{cas}").get_data(as_text=True)
    assert "1 rejected" in client.get("/").get_data(as_text=True)
    import io
    import openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(client.get("/download/status.xlsx").data))["Status"]
    head = [c.value for c in ws[1]]
    row = next(r for r in ws.iter_rows(min_row=2, values_only=True)
               if r[head.index("cas_id")] == cas)
    assert row[head.index("group")] == "Systems"
    assert row[head.index("reviewer")] == "Zhiwen Fan"
    assert row[head.index("first_call")] == "accept"
    assert row[head.index("final_call")] == "reject"


def test_split_evenly_and_membership_changes(office):
    appmod, client, store, ids = office
    _as(client, "office")
    client.post("/groups/add", data={"name": "Systems", "leaders": ["gratz-paul"],
                                     "members": ["fan-zhiwen", "duffield-nick"]})
    for cas in ids:
        client.post("/triage/assign", data={"cas_id": cas, "group": "systems"})

    _as(client, "pgratz")
    client.post("/group/systems/split")
    who = [store.get(cas).get("assigned_to") for cas in ids]
    assert sorted(who) == ["duffield-nick", "duffield-nick", "fan-zhiwen", "fan-zhiwen"]

    # dropping a member sends their open reviews back to the leader
    _as(client, "office")
    client.post("/groups/update", data={"id": "systems", "name": "Systems",
                                        "leaders": ["gratz-paul"], "members": ["fan-zhiwen"]})
    who = [store.get(cas).get("assigned_to") for cas in ids]
    assert who.count("duffield-nick") == 0 and who.count(None) == 2

    # deleting the group sends everyone back to the shortlist, not sent
    client.post("/groups/delete", data={"id": "systems"})
    assert all("group" not in store.get(cas) and "assigned_to" not in store.get(cas)
               for cas in ids)
    assert "Not sent to a group yet" in client.get("/").get_data(as_text=True)


def test_groups_pages_are_the_offices(office):
    appmod, client, store, ids = office
    _as(client, "pgratz")
    for url in ("/groups", "/groups?edit=1"):
        assert client.get(url).status_code == 403
    assert client.post("/groups/add", data={"name": "Mine"}).status_code == 403
    assert client.get("/group/nope").status_code == 404


def test_no_logins_means_one_operator_does_everything(tmp_path, monkeypatch):
    from src import app as appmod
    from src.run import run
    from src.store import Store, db_path_for
    monkeypatch.setattr(appmod, "DATA", str(tmp_path))
    appmod.app.config.update(TESTING=True)
    client = appmod.app.test_client()
    run(CSV, str(tmp_path), FIELD_MAP)
    G.save_groups([SYSTEMS])
    store = Store(db_path_for(str(tmp_path)))
    cas = store.index()[0]["cas_id"]
    store.update(cas, lambda u: {**u, "triage": "shortlist"})
    client.post("/triage/assign", data={"cas_id": cas, "group": "systems"})
    client.post("/group/systems/assign", data={"cas_id": cas, "reviewer": "gratz-paul"})
    # the leader reviewing it themselves: one call counts as both
    client.post("/group/systems/final", data={"cas_id": cas, "decision": "accept"})
    rec = store.get(cas)
    assert rec["professor_status"] == rec["final_status"] == "accept"
    assert client.get("/groups").status_code == 200
    assert client.get("/group/systems").status_code == 200
