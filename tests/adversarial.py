"""Adversarial pass over the grad app: hostile input, injection, abuse, races.

    python tests/adversarial.py [scratch-dir]

Separate from the pytest suite on purpose: it writes junk files, hammers the
store from a dozen threads and pokes the CSRF layer, none of which belongs in a
run you do on every change. Point it at a scratch directory, never at `data/`.

Prints PASS/FAIL per probe; FAIL means the app did something worse than
refusing politely. It has already earned its keep three times:

  * a 500-character filename raised OSError 36 and 500'd the upload,
  * "rejected" in the status column read as undecided,
  * a transcript legend put "Undergraduate degree credit" in a degree column.
"""
import io
import os
import sys
import threading
import time
import zipfile

sys.path.insert(0, ".")

RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((ok, name, detail))
    print(("  PASS  " if ok else "  FAIL  ") + name + (f"  — {detail}" if detail else ""))


def section(t):
    print("\n" + t)
    print("-" * len(t))


def client(tmp, testing=True):
    from src import app as appmod
    from src import users as usersmod
    # Never the real ones. This file's own probes are what taught us why: the
    # XSS and YAML-injection payloads below were landing in the developer's
    # live config/professors.yaml and appearing on the real dashboard.
    from src import triage as triagemod
    usersmod.USERS_YAML = os.path.join(tmp, "users.yaml")
    triagemod.PROFESSORS_YAML = os.path.join(tmp, "professors.yaml")
    appmod.DATA = tmp
    os.makedirs(os.path.join(tmp, "raw"), exist_ok=True)
    appmod.app.config.update(TESTING=testing)
    appmod.app.testing = testing
    return appmod, appmod.app.test_client()


def xlsx(rows, headers):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Applicants"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main(tmp):
    from src.run import run
    from src.store import Store, db_path_for
    appmod, c = client(tmp)
    FIELD_MAP = os.path.join("config", "csv_field_map.yaml")

    # ---------------------------------------------------------------- uploads
    section("Hostile uploads")

    r = c.post("/process", data={"csv": (io.BytesIO(b""), "empty.csv")},
               content_type="multipart/form-data", follow_redirects=True)
    check("empty CSV doesn't 500", r.status_code == 200, f"status {r.status_code}")

    r = c.post("/process", data={"csv": (io.BytesIO(b"MZ\x90\x00 evil"), "payload.csv")},
               content_type="multipart/form-data", follow_redirects=True)
    check("renamed executable refused", b"Refused" in r.data or b"refused" in r.data)

    r = c.post("/process", data={"csv": (io.BytesIO(b"cas_id\n1\n"), "../../../etc/passwd.csv")},
               content_type="multipart/form-data", follow_redirects=True)
    escaped = [p for p in os.listdir(tmp) if "passwd" in p]
    outside = os.path.exists(os.path.join(os.path.dirname(tmp), "etc"))
    check("path traversal in filename contained", not outside, f"raw has {escaped}")

    r = c.post("/process", data={"csv": (io.BytesIO(b"cas_id\n1\n"), "x" * 500 + ".csv")},
               content_type="multipart/form-data", follow_redirects=True)
    check("500-char filename survives", r.status_code == 200, f"status {r.status_code}")

    # a zip whose member escapes the extraction dir (zip-slip)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("../../../../tmp/zipslip_1000419138.pdf", b"%PDF-1.4 not really")
    r = c.post("/process", data={"zip": (io.BytesIO(buf.getvalue()), "slip.zip")},
               content_type="multipart/form-data", follow_redirects=True)
    check("zip-slip writes nothing outside", not os.path.exists("/tmp/zipslip_1000419138.pdf"))

    # a PDF that is not a PDF
    r = c.post("/process", data={"zip": (io.BytesIO(b"<html>gotcha</html>"),
                                         "1000419138_x_Full Application_1.pdf")},
               content_type="multipart/form-data", follow_redirects=True)
    check("fake PDF refused, no crash", r.status_code == 200)

    # non-UTF8 CSV
    latin = "cas_id,first_name,last_name\n1000111222,José,Muñoz\n".encode("latin-1")
    r = c.post("/process", data={"csv": (io.BytesIO(latin), "latin.csv")},
               content_type="multipart/form-data", follow_redirects=True)
    check("latin-1 CSV doesn't 500", r.status_code == 200, f"status {r.status_code}")

    # headers only, no rows
    r = c.post("/process", data={"csv": (io.BytesIO(b"cas_id,first_name\n"), "hdr.csv")},
               content_type="multipart/form-data", follow_redirects=True)
    check("header-only CSV handled", r.status_code == 200)

    # ------------------------------------------------------------- injection
    section("Injection")

    nasty = '<script>alert(1)</script>'
    csv = f"cas_id,first_name,last_name\n1000999001,{nasty},Bobby\n".encode()
    c.post("/process", data={"csv": (io.BytesIO(csv), "xss.csv")},
           content_type="multipart/form-data", follow_redirects=True)
    page = c.get("/applicants").get_data(as_text=True)
    check("applicant name is escaped in HTML",
          "<script>alert(1)</script>" not in page and "&lt;script&gt;" in page)

    store = Store(db_path_for(tmp))
    ids = [e["cas_id"] for e in store.index()]
    if ids:
        store.update(ids[0], lambda u: {**u, "review_status": "tbd"})
        deck_page = c.get("/triage").get_data(as_text=True)
        check("deck escapes the same name", "<script>alert(1)</script>" not in deck_page)

    r = c.post("/professors/add", data={"name": "<img src=x onerror=alert(1)>",
                                        "url": "https://x.test/p"}, follow_redirects=True)
    page = c.get("/professors").get_data(as_text=True)
    check("professor name escaped", "<img src=x onerror" not in page)
    r = c.post("/professors/add", data={"name": "JS Guy", "url": "javascript:alert(1)"},
               follow_redirects=True)
    check("javascript: profile URL refused", b"http://" in r.data or b"https://" in r.data)

    # YAML round-trip safety: a name full of YAML metacharacters
    weird = "*anchor: &a [{'x': 1}]\n#comment"
    c.post("/professors/add", data={"name": weird}, follow_redirects=True)
    from src.triage import load_professors
    check("YAML-metacharacter name survives a reload",
          any(p["name"] == weird for p in load_professors()),
          f"{len(load_professors())} professors on file")

    # spreadsheet formula in the status cell (CSV/Excel injection)
    data = xlsx([["1000999001", "=cmd|'/C calc'!A0"]], ["cas_id", "status"])
    r = c.post("/step/3", data={"csv": (io.BytesIO(data), "formula.xlsx")},
               content_type="multipart/form-data", follow_redirects=True)
    rec = Store(db_path_for(tmp)).get("1000999001") or {}
    check("formula in status is not treated as a decision",
          rec.get("review_status") in ("", "tbd", None) or rec.get("review_status") not in ("accept", "reject"),
          f"status={rec.get('review_status')!r}")

    # ------------------------------------------------------------ status quirks
    section("Status column quirks")

    for value, expect in [("  ACCEPT  ", "accept"), ("Tbd.", "tbd"), ("rejected", "reject"),
                          ("¯\\_(ツ)_/¯", ""), ("accept reject", ""), ("0", ""),
                          ("nan", ""), ("None", "")]:
        from src.review import normalize_status
        got = normalize_status(value)
        check(f"status {value!r} -> {got!r}", got == expect, f"expected {expect!r}")

    # a status column named oddly
    from src.review import status_from_row
    check("' Status ' header still read",
          status_from_row({" Status ": "tbd"}) == "tbd")
    check("'STATUS' header still read", status_from_row({"STATUS": "accept"}) == "accept")
    check("'local_status_0' is NOT mistaken for the review column",
          status_from_row({"local_status_0": "Final ADMITS for ADMI"}) == "")

    # ------------------------------------------------------------------ routes
    section("Route abuse")

    check("unknown student 404s", c.get("/student/does-not-exist").status_code == 404)
    check("traversal in cas_id 404s",
          c.get("/student/..%2f..%2fetc%2fpasswd").status_code in (404, 400, 308))
    check("unknown job 404s", c.get("/jobs/deadbeef").status_code == 404)
    check("swipe on unknown id 404s",
          c.post("/triage/swipe", data={"cas_id": "nope", "direction": "reject"}).status_code == 404)
    check("bogus direction 400s",
          c.post("/triage/swipe", data={"cas_id": ids[0] if ids else "x",
                                        "direction": "../../etc"}).status_code == 400)
    check("assign to unknown professor 400s",
          c.post("/triage/assign", data={"cas_id": ids[0] if ids else "x",
                                         "professor": "ghost"}).status_code == 400)
    check("delete unknown professor doesn't 500",
          c.post("/professors/delete", data={"id": "ghost"},
                 follow_redirects=True).status_code == 200)
    check("GET on a POST-only route is refused",
          c.get("/triage/swipe").status_code in (405, 404))

    # ----------------------------------------------------------------- CSRF
    section("CSRF and same-origin")
    appmod2, c2 = client(tmp, testing=False)
    r = c2.post("/reset", data={})
    check("POST without a CSRF token is rejected", r.status_code == 400, f"status {r.status_code}")
    r = c2.post("/reset", data={}, headers={"Origin": "https://evil.example"})
    check("cross-origin POST is rejected", r.status_code == 403, f"status {r.status_code}")
    appmod2.app.config.update(TESTING=True)
    appmod2.app.testing = True

    # ------------------------------------------------------------ concurrency
    section("Concurrency")

    run("tests/fixtures/sample_applications.csv", tmp, FIELD_MAP)
    store = Store(db_path_for(tmp))
    ids = [e["cas_id"] for e in store.index()][:4]
    for sid in ids:
        store.update(sid, lambda u: {**u, "review_status": "tbd"})

    errors = []

    def hammer(sid, direction):
        try:
            for _ in range(12):
                c.post("/triage/swipe", data={"cas_id": sid, "direction": direction})
        except Exception as exc:      # noqa: BLE001
            errors.append(str(exc))

    threads = [threading.Thread(target=hammer, args=(sid, d))
               for sid in ids for d in ("shortlist", "reject", "undo")]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    check("40+ concurrent swipes don't corrupt the store", not errors, "; ".join(errors[:2]))
    after = Store(db_path_for(tmp)).index()
    check("every record still readable after the race", len(after) >= len(ids),
          f"{len(after)} records")

    # two uploads at once
    def upload():
        with open("tests/fixtures/sample_applications.csv", "rb") as f:
            c.post("/process", data={"csv": (io.BytesIO(f.read()), "race.csv")},
                   content_type="multipart/form-data", follow_redirects=True)

    ts = [threading.Thread(target=upload) for _ in range(4)]
    for t in ts:
        t.start()
    for t in ts:
        t.join()
    n = Store(db_path_for(tmp)).count()
    check("four simultaneous uploads don't duplicate rows", n == len(after),
          f"{n} records, was {len(after)}")

    # ------------------------------------------------------------------ jobs
    section("Background jobs")
    from src import jobs

    jid = jobs.start(lambda p: 1 / 0, "boom")
    for _ in range(100):
        j = jobs.get(jid)
        if j["done"]:
            break
        time.sleep(0.02)
    check("a job that raises is reported, not swallowed", bool(j.get("error")),
          str(j.get("error"))[:60])
    check("failed job's status endpoint still answers",
          c.get(f"/jobs/{jid}").status_code == 200)

    for i in range(60):
        jobs.start(lambda p: None, f"j{i}")
    time.sleep(0.4)
    check("job registry is capped", len(jobs._JOBS) <= 41, f"{len(jobs._JOBS)} jobs held")

    # ------------------------------------------------------- degree extraction
    section("Degree extraction traps")
    from src.review import undergrad_degree

    def rec_with(text, key="resume"):
        return {"sources": {key: {"text": text, "analysis": {}}}}

    traps = [
        ("...and be a part of a renowned university", "", "the word 'be'"),
        ("NDC = Not for Undergraduate degree credit", "", "transcript boilerplate"),
        ("Master of Science in Robotics", "", "graduate only"),
        ("She labs about bs all day", "", "'bs' inside words"),
        ("Bachelor of Science in Photonics", "Bachelor of Science in Photonics", "the real thing"),
    ]
    for text, expect, why in traps:
        got = undergrad_degree(rec_with(text))
        check(f"degree from {why!r}", got == expect, f"got {got!r}, wanted {expect!r}")

    # ---------------------------------------------------------- review sheet
    section("Review sheet edge cases")
    from src.review import write_review_sheet
    import openpyxl

    out = os.path.join(tmp, "empty_review.xlsx")
    try:
        write_review_sheet([], os.path.join(tmp, "raw"), out)
        ws = openpyxl.load_workbook(out)["Applicants"]
        check("review sheet with zero applicants still opens", ws.max_row >= 1,
              f"{ws.max_row} row(s)")
    except Exception as exc:                       # noqa: BLE001
        check("review sheet with zero applicants still opens", False, str(exc)[:80])

    recs = Store(db_path_for(tmp)).index()
    out2 = os.path.join(tmp, "twice_review.xlsx")
    write_review_sheet(recs, os.path.join(tmp, "raw"), out2)
    hdr1 = [c_.value for c_ in openpyxl.load_workbook(out2)["Applicants"][1]]
    write_review_sheet(recs, os.path.join(tmp, "raw"), out2)
    hdr2 = [c_.value for c_ in openpyxl.load_workbook(out2)["Applicants"][1]]
    check("re-exporting doesn't stack duplicate columns", hdr1 == hdr2,
          f"{hdr1.count('status')} status columns")

    # ------------------------------------------------- self-signup abuse
    section("Self-signup")
    from src import users as usersmod
    signup_dir = os.path.join(tmp, "signup")
    appmod3, c3 = client(signup_dir)
    usersmod.save_users(usersmod.add_user([], "boss", "pw1234", usersmod.ADVISOR))

    # the obvious one: ask for the role you want in the form
    c3.post("/signup", data={"username": "climber", "password": "pw1234",
                             "confirm": "pw1234", "role": "advisor",
                             "professor": "gratz-paul"})
    row = next((u for u in usersmod.load_users() if u["username"] == "climber"), None)
    check("signup can't grant itself a role", row is not None and row["role"] == "pending",
          f"role={row and row['role']}")
    check("signup can't attach itself to a professor",
          row is not None and row["professor"] == "", f"professor={row and row['professor']}")

    r = c3.post("/login", data={"username": "climber", "password": "pw1234"},
                follow_redirects=True)
    check("a pending account can't sign in", b"waiting for an advisor" in r.data)
    for gated in ("/applicants", "/accounts", "/prof", "/download/status.xlsx"):
        rr = c3.get(gated)
        check(f"pending account is refused {gated}", rr.status_code in (302, 401, 403),
              f"status {rr.status_code}")

    # approving is advisor-only, and can't invent a role
    rr = c3.post("/accounts/approve", data={"username": "climber", "role": "advisor"})
    check("approve refused to a signed-out visitor", rr.status_code in (302, 401, 403),
          f"status {rr.status_code}")
    c3.post("/login", data={"username": "boss", "password": "pw1234"})
    c3.post("/accounts/approve", data={"username": "climber", "role": "superuser"})
    row = next(u for u in usersmod.load_users() if u["username"] == "climber")
    check("approve rejects a role outside the two real ones", row["role"] == "pending",
          f"role={row['role']}")

    # a request can't be used to overwrite an existing account's password
    c3.get("/logout")
    c3.post("/signup", data={"username": "boss", "password": "hijacked",
                             "confirm": "hijacked"})
    check("signup can't overwrite an existing account",
          usersmod.authenticate("boss", "hijacked") is None
          and usersmod.authenticate("boss", "pw1234") is not None)
    check("duplicate signup leaves exactly one record for that name",
          sum(1 for u in usersmod.load_users() if u["username"] == "boss") == 1)

    # a lone pending request must not flip a no-login app into locked-out
    empty_dir = os.path.join(tmp, "signup_empty")
    appmod4, c4 = client(empty_dir)
    usersmod.save_users(usersmod.request_account([], "ghost", "pw1234"))
    check("a pending request alone doesn't turn login on",
          not usersmod.login_required_globally() and c4.get("/").status_code == 200)
    check("the advisor can still reach setup afterwards",
          c4.get("/setup").status_code == 200)

    print("\n" + "=" * 62)
    bad = [r for r in RESULTS if not r[0]]
    print(f"{len(RESULTS) - len(bad)} passed, {len(bad)} failed")
    for _, name, detail in bad:
        print(f"  FAILED: {name}  — {detail}")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1 else "/tmp/advdata"))
