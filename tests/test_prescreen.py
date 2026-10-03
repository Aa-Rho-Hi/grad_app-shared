"""Prescreen rules (src/screening.py) and the spreadsheets they feed
(src/prescreen_report.py): the grad office's Quick Admit / MS Distance / PhD
Distance rules, applied as flags. Synthetic applicants only."""
import csv
import json
import os
import sys
import zipfile

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import yaml

from src import screening as S
from src.normalize import build_record
from src.models import parse_record

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CFG = yaml.safe_load(open(os.path.join(ROOT, "config", "csv_field_map.yaml")))
RULES = S.load_rules()
APPROVED = S.load_courses()
LETTERS = RULES["letters"]

QA = "Fall 2027 Quick Admit M.S. Electrical Engineering"
MS_NT = "Fall 2027 M.S. Electrical Engineering - Distance Non-Thesis"
MS_T = "Fall 2027 M.S. Electrical Engineering - Distance Thesis"
PHD = "Fall 2027 Ph.D. Electrical Engineering - Distance"
TEN = ["ECEN 444", "ECEN 474", "ECEN 601", "ECEN 602", "ECEN 605", "ECEN 607",
       "ECEN 608", "ECEN 609", "ECEN 610", "ECEN 613"]


def _letter(key, name):
    return LETTERS[key].replace("[NAME]", name)


# ------------------------------------------------------------- classification

@pytest.mark.parametrize("name,track,thesis", [
    (QA, S.QUICK_ADMIT, None),
    (MS_NT, S.MS_DISTANCE, False),
    (MS_T, S.MS_DISTANCE, True),
    ("Spring 2028 M.S. Computer Engineering (Online) - Thesis", S.MS_DISTANCE, True),
    (PHD, S.PHD_DISTANCE, None),
    ("Fall 2027 Doctor of Philosophy Computer Engineering - Distance", S.PHD_DISTANCE, None),
    ("Fall 2027 M.Eng. Electrical Engineering - Distance", S.OTHER_DISTANCE, None),
    ("Fall 2020 M.S. Electrical Engineering (ECEN) - International Students", S.ON_CAMPUS, None),
    ("Spring 2021 M.S. Non-Thesis Electrical Engineering (ECEN)", S.ON_CAMPUS, False),
])
def test_track_comes_from_the_designation(name, track, thesis):
    got = S.classify([{"index": 0, "name": name}], RULES)
    assert got["track"] == track and got["thesis"] == thesis


def test_quick_admit_distance_follows_the_distance_rules():
    # The Quick Admit rules are "on campus programs only".
    got = S.classify([{"index": 0, "name": "Quick Admit M.S. EE - Distance Non-Thesis"}], RULES)
    assert got["track"] == S.MS_DISTANCE


# -------------------------------------------------------------------- courses

def test_course_codes_normalize():
    text = "Plan: ECE 601, ECEN-602; ecen603 and ECEN 7041 (not a course) CSCE 629"
    assert S.extract_courses(text) == ["ECEN 601", "ECEN 602", "ECEN 603", "CSCE 629"]


def test_ten_approved_courses_with_two_400_level_pass():
    c = S.check_courses(TEN, 10, APPROVED, RULES["courses"])
    assert c["ok"] and c["counted"] == 10


def test_a_third_400_level_course_fails_the_cap():
    listed = TEN + ["ECEN 420"]
    c = S.check_courses(listed, 10, APPROVED, RULES["courses"])
    assert not c["ok"]
    assert any("3 400-level" in p for p in c["problems"])


def test_off_list_and_low_level_courses_do_not_count():
    listed = TEN[:-2] + ["ECEN 689", "ECEN 303"]          # 689 not on list, 303 too low
    c = S.check_courses(listed, 10, APPROVED, RULES["courses"])
    assert c["counted"] == 8 and not c["ok"]
    assert c["not_on_list"] == ["ECEN 689"] and c["below_400"] == ["ECEN 303"]
    assert any("8 approved course(s) listed; 10 needed" in p for p in c["problems"])


def test_course_list_prefers_application_answers_and_skips_transcripts():
    app = {"course_answers": ["ECEN 601, ECEN 602"]}
    pages = [{"page": 3, "label": "scores", "text": "Course Plan ECEN 700 ECEN 710 ECEN 720"}]
    assert S.find_course_list(app, pages, ["ECEN"])[0] == ["ECEN 601", "ECEN 602"]
    transcript = {"page": 2, "label": "scores",
                  "text": "Official Transcript\nCredit Hours  Quality Points\n"
                          "ECEN 314 A\nECEN 325 B\nECEN 350 A"}
    plan = {"page": 5, "label": "ignore",
            "text": "Distance learning course plan\nECEN 601\nECEN 602\nECEN 605"}
    courses, where = S.find_course_list({}, [transcript, plan], ["ECEN"])
    assert courses == ["ECEN 601", "ECEN 602", "ECEN 605"] and where == "packet p.5"


# -------------------------------------------------------------------- letters

def test_exact_wording_survives_pdf_mangling():
    text = _letter("ms_employed", "Jane Doe")
    # what a PDF text layer does to a letter: odd spacing, curly quotes,
    # line breaks mid-sentence, CAS's sprinkled spaces
    mangled = ("ACME CORP\n" + text.replace("We give approval", "W e give  approval")
               .replace(". ", ".\n").replace("MS research", "M S research")
               .replace("on-campus", "on‑campus") + "\nSincerely, Pat Manager")
    got = S.match_letter([{"page": 7, "text": mangled}], LETTERS["ms_employed"],
                         LETTERS, "Jane Doe")
    assert got["status"] == "exact" and got["where"] == "p.7"
    assert "name_mismatch" not in got


def test_letter_split_across_a_page_break_is_found():
    text = _letter("phd_employed", "Jane Doe")
    half = len(text) // 2
    pages = [{"file": "f", "page": 4, "text": text[:half]},
             {"file": "f", "page": 5, "text": text[half:]}]
    got = S.match_letter(pages, LETTERS["phd_employed"], LETTERS, "Jane Doe")
    assert got["status"] == "exact" and got["where"] == "p.4-5"


def test_changed_wording_is_close_not_exact():
    text = _letter("ms_employed", "Jane Doe").replace("disseminate", "share")
    got = S.match_letter([{"page": 2, "text": text}], LETTERS["ms_employed"], LETTERS, "Jane Doe")
    assert got["status"] == "close" and 0.85 <= got["coverage"] < 1


def test_unrelated_page_is_missing():
    got = S.match_letter([{"page": 1, "text": "I am writing to recommend the applicant "
                                              "to the program. " * 20}],
                         LETTERS["ms_employed"], LETTERS)
    assert got["status"] == "missing"


def test_assess_typo_in_the_official_text_matches_access():
    # The PhD wording the office sent says "on-campus assess"; an applicant
    # copying the correct word must not fail on it.
    tpl = LETTERS["phd_employed"].replace("access", "assess")
    got = S.match_letter([{"page": 1, "text": _letter("phd_employed", "Jane Doe")}],
                         tpl, LETTERS, "Jane Doe")
    assert got["status"] == "exact"


def test_letter_for_someone_else_is_flagged():
    got = S.match_letter([{"page": 1, "text": _letter("ms_employed", "John Roe")}],
                         LETTERS["ms_employed"], LETTERS, "Jane Doe")
    assert got["status"] == "exact" and "johnroe" in got["name_mismatch"]


def test_advisor_proof_found_in_a_letter_not_in_the_sop():
    sop = {"page": 2, "label": "sop", "text": "Prof. Smith has agreed to serve as my advisor."}
    assert S.find_advisor_proof([sop])["status"] == "missing"
    letter = {"page": 9, "label": "lor",
              "text": "Dear Graduate Office,\nI am happy to serve as the research advisor "
                      "for Jane Doe for her M.S. thesis.\nRegards, Prof. Smith"}
    got = S.find_advisor_proof([sop, letter])
    assert got["status"] == "found" and got["where"] == "p.9"


def test_declining_to_advise_is_not_proof():
    page = {"page": 3, "label": "lor",
            "text": "Unfortunately I am not willing to serve as her advisor this year."}
    assert S.find_advisor_proof([page])["status"] == "missing"


# --------------------------------------------------------------- GPA and dept

def _unified(app, resume_text=None):
    u = {"student_id": "1", "identity": {"full_name": "Jane Doe"},
         "sources": {"application": app, "resume": None, "sop": None, "scores": None,
                     "lors": []}, "warnings": []}
    if resume_text:
        u["sources"]["resume"] = {"text": resume_text}
    return u


def _edu(name, raw, scale=4.0, **kw):
    return dict({"index": 0, "college_name": name,
                 "gpa": {"raw": raw, "scale": scale, "normalized_4": min(raw / scale * 4, 4)}}, **kw)


@pytest.mark.parametrize("raw,status", [(3.25, "pass"), (3.249, "fail"), (3.9, "pass")])
def test_gpa_threshold_edges(raw, status):
    u = _unified({"education": [_edu("Somewhere", raw)]})
    assert S.gpa_check(u, S.MS_DISTANCE, 3.25)["status"] == status


def test_cas_cumulative_gpa_wins():
    u = _unified({"education": [_edu("A", 2.9)], "gpa_summary": {"undergraduate": 3.4}})
    assert S.gpa_check(u, S.QUICK_ADMIT, 3.25)["status"] == "pass"


def test_quick_admit_uses_the_texas_am_gpa():
    u = _unified({"education": [_edu("Blinn College", 2.8),
                                _edu("Texas A&M University", 3.6)]})
    assert S.gpa_check(u, S.QUICK_ADMIT, 3.25)["status"] == "pass"
    assert S.gpa_check(u, S.MS_DISTANCE, 3.25)["status"] == "mixed"


def test_gpa_on_another_scale_is_not_converted():
    u = _unified({"education": [_edu("IIT", 9.1, scale=10.0)]})
    got = S.gpa_check(u, S.MS_DISTANCE, 3.25)
    assert got["status"] == "scale" and "9.1/10" in got["text"]


@pytest.mark.parametrize("major,status", [
    ("Electrical Engineering", "ece"), ("Computer Engineering", "ece"),
    ("Electrical & Computer Engineering", "ece"),
    ("Electrical Engineering Technology", "other"),
    ("Mechanical Engineering", "other"), ("Computer Science", "other"),
])
def test_ece_department(major, status):
    u = _unified({"background": major})
    assert S.ece_check(u, RULES)["status"] == status


def test_major_unknown_goes_to_review():
    u = _unified({"programs": [{"index": 0, "name": QA}], "education": [_edu("TAMU", 3.8)]})
    got = S.evaluate(u, RULES, APPROVED)
    assert got["outcome"] == S.REVIEW
    assert "couldn't find the applicant's major/department" in got["reasons"]


def test_masters_detection():
    assert S.masters_check(_unified({"education": [
        {"index": 0, "degree": "Bachelor of Science"},
        {"index": 1, "degree": "Master of Science"}]}))["has_ms"] is True
    assert S.masters_check(_unified({"education": [
        {"index": 0, "degree": "Bachelor of Science"}]}))["has_ms"] is False
    assert S.masters_check(_unified({}, "M.S. Electrical Engineering, expected May 2027"))[
        "has_ms"] is None
    assert S.masters_check(_unified({}, "Skills: MS Office, Python\nB.Tech EE 2019"))[
        "has_ms"] is False


# ---------------------------------------------------- normalize / privacy

def test_new_roster_fields_and_residency_only_label():
    row = {"cas_id": "77", "first_name": "A", "last_name": "B",
           "designation_0": "Fall 2027 M.S. EE - Distance Thesis",
           "college_name_0": "Texas A&M University", "gpas_by_transcript_gpa_0": "3.5",
           "college_degree_0": "Bachelor of Science", "college_major_0": "Electrical Engineering",
           "gre_official_verbal_converted": "155", "gre_official_quantitative_converted": "165",
           "toefl_ibt_official_score_0": "", "toefl_ibt_result_0": "91",
           "duolingo_english_test_overall_score_0": "125", "duolingo_literacy_0": "140",
           "honors_and_awards": "Dean's List", "custom_questions_planned_courses": "ECEN 601",
           "citizenship_status": "U.S. Permanent Resident", "citizenship": "Freedonia",
           "gpas_by_transcript_cumulative_undergraduate_gpa": "0"}
    rec, errors = parse_record(build_record(row, CFG, "x.csv", 0))
    assert not errors
    assert rec["education"][0]["degree"] == "Bachelor of Science"
    assert rec["education"][0]["major"] == "Electrical Engineering"
    assert rec["gre"] == {"verbal": 155.0, "quant": 165.0, "official": True}
    eng = rec["english_proficiency"]
    assert eng["best_toefl_total"] == 91.0          # an empty official column used to hide this
    assert eng["best_duolingo"] == 125.0            # the overall, not the higher subscore
    assert rec["awards"] == ["Dean's List"]
    assert rec["course_answers"] == ["ECEN 601"]
    assert rec["residency"] == "Domestic"
    assert "gpa_summary" not in rec                 # CAS's 0 means "not calculated"
    assert "Freedonia" not in json.dumps(rec)       # the country never enters a record


# ----------------------------------------------- end to end: roster + packets

def _pdf(path, pages):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(str(path), pagesize=LETTER)
    for text in pages:
        y = 740
        c.setFont("Helvetica", 10)
        for para in text.split("\n"):
            words, line = para.split(), ""
            for w in words + [None]:
                if w is None or len(line) + len(w) > 95:
                    c.drawString(60, y, line)
                    y -= 14
                    line = ""
                if w:
                    line = (line + " " + w).strip()
        c.showPage()
    c.save()


def _plan(courses):
    return "ECEN Distance Learning Course Plan\nPlanned courses:\n" + "\n".join(courses)


ADVISOR = ("Dear Graduate Admissions Committee,\nI have agreed to serve as the research "
           "advisor for {name} during the M.S. thesis.\nRegards,\nProf. Alex Smith\n"
           "Department of Electrical and Computer Engineering, Texas A&M University")
SOP = ("Statement of Purpose\n" + "I am applying to the program because my research "
       "interest is power systems and I plan to pursue graduate study. " * 6)


@pytest.fixture
def cohort(tmp_path):
    pytest.importorskip("reportlab")
    from src.ingest_zip import ingest_zip
    from src.run import run
    from src.store import Store, db_path_for

    people = [
        # cas, first, last, designation, gpa, background, extra roster cols, packet pages
        ("1001", "Ana", "Quick", QA, "3.6", "Electrical Engineering", {}, None),
        ("1002", "Ben", "Low", QA, "3.1", "Computer Engineering", {}, None),
        ("1003", "Cal", "Meen", QA, "3.9", "Mechanical Engineering", {}, None),
        ("1004", "Dee", "Nonthesis", MS_NT, "3.5", "Electrical Engineering", {},
         [SOP, _plan(TEN)]),
        ("1005", "Eve", "Thesis", MS_T, "3.5", "Electrical Engineering", {},
         [SOP, _plan(TEN), ADVISOR.format(name="Eve Thesis"),
          "ACME Corp\n" + _letter("ms_employed", "Eve Thesis") + "\nSincerely, HR"]),
        ("1006", "Fay", "Noletter", MS_T, "3.5", "Electrical Engineering", {},
         [SOP, _plan(TEN), ADVISOR.format(name="Fay Noletter")]),
        ("1007", "Gus", "Phd", PHD, "3.0", "Electrical Engineering",
         {"college_degree_0": "Bachelor of Science", "college_degree_1": "Master of Science"},
         [SOP, _plan(TEN[2:8]), _letter("phd_employed", "Gus Phd")]),
        ("1008", "Hal", "Bach", PHD, "3.0", "Electrical Engineering",
         {"college_degree_0": "Bachelor of Science"},
         [SOP, _plan(TEN[2:8]), _letter("phd_employed", "Hal Bach")]),
        ("1009", "Ivy", "Unemployed", MS_T, "3.5", "Electrical Engineering",
         {"custom_questions_courses_planned": ", ".join(TEN)},
         [SOP, ADVISOR.format(name="Ivy Unemployed"), _letter("ms_unemployed", "Ivy Unemployed")]),
        ("1010", "Jo", "Campus", "Fall 2027 M.S. Electrical Engineering", "3.9",
         "Electrical Engineering", {}, None),
    ]
    extra_cols = sorted({k for p in people for k in p[6]})
    cols = ["cas_id", "first_name", "last_name", "email", "designation_0",
            "college_name_0", "gpas_by_transcript_gpa_0", "custom_field_background",
            "citizenship_status"] + extra_cols
    roster = tmp_path / "roster.csv"
    with open(roster, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for cas, first, last, desig, gpa, bg, extra, _pages in people:
            w.writerow([cas, first, last, f"{first.lower()}@x.edu", desig,
                        "Texas A&M University", gpa, bg, "U.S. Citizen"]
                       + [extra.get(c, "") for c in extra_cols])
    zpath = tmp_path / "packets.zip"
    with zipfile.ZipFile(zpath, "w") as z:
        for cas, first, last, desig, _g, _b, _e, pages in people:
            if pages:
                pdf = tmp_path / f"{cas}_{first}_{last}_Full Application_1_Fall 2027.pdf"
                _pdf(pdf, pages)
                z.write(pdf, pdf.name)
    out = tmp_path / "data"
    run(str(roster), str(out), os.path.join(ROOT, "config", "csv_field_map.yaml"))
    ingest_zip(str(zpath), str(out))
    entries = {e["cas_id"]: e for e in Store(db_path_for(str(out))).index()}
    return entries, out


def test_quick_admit_rules(cohort):
    e, _ = cohort
    assert e["1001"]["screening"]["outcome"] == S.ADMIT
    assert e["1002"]["screening"]["outcome"] == S.REVIEW
    assert any("below 3.25" in r for r in e["1002"]["screening"]["reasons"])
    assert e["1003"]["screening"]["outcome"] == S.REVIEW
    assert any("not from ECE" in r for r in e["1003"]["screening"]["reasons"])
    assert e["1010"]["screening"]["outcome"] == ""          # no rule for regular on-campus


def test_ms_distance_rules(cohort):
    e, _ = cohort
    assert e["1004"]["screening"]["outcome"] == S.ADMIT, e["1004"]["screening"]["reasons"]
    eve = e["1005"]["screening"]
    assert eve["outcome"] == S.ADMIT, eve["reasons"]
    assert eve["checks"]["employer_letter"]["status"] == "exact"
    assert any("letterhead" in v for v in eve["verify"])     # a person still looks
    fay = e["1006"]["screening"]
    assert fay["outcome"] == S.REVIEW
    assert "employer letter / unemployment statement not found" in fay["reasons"]
    ivy = e["1009"]["screening"]
    assert ivy["outcome"] == S.ADMIT, ivy["reasons"]
    assert ivy["checks"]["employer_letter"]["option"].startswith("B")
    assert ivy["checks"]["courses"]["source"] == "application course question"


def test_phd_distance_rules(cohort):
    e, _ = cohort
    gus = e["1007"]["screening"]
    assert gus["outcome"] == S.ADMIT, gus["reasons"]
    assert gus["checks"]["courses"]["required"] == 6
    hal = e["1008"]["screening"]
    assert hal["outcome"] == S.REVIEW
    assert any("6 approved course(s) listed; 14 needed" in r for r in hal["reasons"])


def test_spreadsheets_split_link_and_never_execute_text(cohort, tmp_path):
    from openpyxl import load_workbook
    from src.prescreen_report import write_all
    entries, _ = cohort
    entries = dict(entries)
    entries["1001"] = dict(entries["1001"], profile=dict(entries["1001"]["profile"],
                                                         awards=["=HYPERLINK(\"http://x\")"]))
    paths = write_all(entries.values(), "https://gradapp.example/", str(tmp_path / "out"))

    main = load_workbook(paths["summary"])["Applicants"]
    head = [c.value for c in main[1]]
    ids = [r[head.index("CAS ID")] for r in main.iter_rows(min_row=2, values_only=True)]
    assert set(ids) == {"1001", "1002", "1003", "1010"}       # distance moved out
    row = next(r for r in main.iter_rows(min_row=2) if r[head.index("CAS ID")].value == "1001")
    assert row[0].hyperlink.target == "https://gradapp.example/student/1001"
    award = row[head.index("Awards")]
    assert award.data_type == "s" and award.value.startswith("=HYPERLINK")

    ms = load_workbook(paths["ms-distance"])["MS Distance"]
    ms_ids = {r[1] for r in ms.iter_rows(min_row=2, values_only=True)}
    assert ms_ids == {"1004", "1005", "1006", "1009"}
    phd = load_workbook(paths["phd-distance"])["PhD Distance"]
    assert {r[1] for r in phd.iter_rows(min_row=2, values_only=True)} == {"1007", "1008"}
    assert {"Summaries", "Rules"} <= set(load_workbook(paths["phd-distance"]).sheetnames)
    summ = load_workbook(paths["summary"])["Summaries"]
    first = next(summ.iter_rows(min_row=2))
    assert first[0].hyperlink.target.startswith("https://gradapp.example/student/")
    assert first[4].value and "Applied:" in first[4].value


def test_prescreen_page_and_downloads(cohort, monkeypatch):
    from src import app as appmod
    _, out = cohort
    monkeypatch.setattr(appmod, "DATA", str(out))
    appmod.app.testing = True
    client = appmod.app.test_client()
    page = client.get("/prescreen")
    assert page.status_code == 200 and b"MS Distance" in page.data
    assert b"Further review" in client.get("/prescreen?outcome=Further+review").data
    xl = client.get("/download/prescreen/ms-distance.xlsx")
    assert xl.status_code == 200 and xl.data[:2] == b"PK"
    z = client.get("/download/prescreen.zip")
    assert z.status_code == 200
    import io
    names = zipfile.ZipFile(io.BytesIO(z.data)).namelist()
    assert sorted(names) == ["prescreen_ms_distance.xlsx", "prescreen_phd_distance.xlsx",
                             "prescreen_summary.xlsx"]
    assert client.get("/download/prescreen/nope.xlsx").status_code == 404
    assert client.post("/prescreen/reapply").status_code == 302
    student = client.get("/student/1005")
    assert b"Prescreen" in student.data and b"Admit" in student.data


def test_review_sheet_gets_a_link_column(cohort, tmp_path):
    from openpyxl import load_workbook
    from src.review import LINK_COLUMN, write_review_sheet
    entries, out = cohort
    path = write_review_sheet(list(entries.values()), str(out / "raw"),
                              str(tmp_path / "review.xlsx"), link_base="http://127.0.0.1:5001")
    ws = load_workbook(path)["Applicants"]
    head = [c.value for c in ws[1]]
    col = head.index(LINK_COLUMN)
    links = {r[head.index("cas_id")].value: r[col].hyperlink.target
             for r in ws.iter_rows(min_row=2) if r[col].hyperlink}
    assert links["1001"] == "http://127.0.0.1:5001/student/1001"


def test_stale_rules_are_recomputed_on_read(cohort, monkeypatch, tmp_path):
    from src.store import Store, db_path_for
    _, out = cohort
    stricter = dict(RULES, gpa_threshold=3.7)
    path = tmp_path / "rules.yaml"
    path.write_text(yaml.safe_dump(stricter))
    monkeypatch.setattr(S, "RULES_YAML", str(path))
    e = {x["cas_id"]: x for x in Store(db_path_for(str(out))).index()}
    assert e["1001"]["screening"]["outcome"] == S.REVIEW     # 3.6 < 3.7 now
